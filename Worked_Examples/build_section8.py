"""Generate Section 8 (Worked Examples) as a standalone .docx."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/Users/kanushetkar/Desktop/NTSB Paper Process/Section 8 Worked Examples.docx"

doc = Document()
# Base style
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)


def h(text, level):
    doc.add_heading(text, level=level)


def p(text="", italic=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.italic = italic
    return par


def kv(label, text):
    par = doc.add_paragraph()
    r = par.add_run(label + " ")
    r.bold = True
    par.add_run(text)
    return par


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, head in enumerate(headers):
        hdr[i].text = head
        for par in hdr[i].paragraphs:
            for run in par.runs:
                run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()
    return t


def eq(text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run(text)
    run.italic = True
    return par


# ===================== SECTION 8 =====================
h("8. Worked Examples", 1)
p("We illustrate the pipeline on two incident types: an engine fire and a landing-gear "
  "malfunction. For each, we run two query styles. The after-incident query is the full "
  "investigator narrative written after the event; because a complete aftermath already "
  "contains the downstream events, we report only diagnosis for it. The during-incident query "
  "is a short, real-time report from the flight crew; for it we report both diagnosis and "
  "prognosis. NTSB probable-cause codes are stored as hierarchical strings (for example, "
  "aircraft-power plant-engine-turbine section-fatigue/wear/corrosion - c); we pass these "
  "through an LLM that rewrites them into plain-English labels for readability, but all "
  "probabilities are computed on the raw codes. For each cluster we report P(K|Q), and for "
  "each cause the final P(C|Q) from the law of total probability.")

# ---------- 8.1 Engine Fire ----------
h("8.1 Engine Fire", 2)

h("8.1.1 After-incident query (diagnosis only)", 3)
kv("Query (investigator narrative):",
   "\"On January 11, 2010, an Aerospatiale ATR72 experienced a No. 1 (left) engine fire "
   "during takeoff. The pilot declared an emergency, shut down the No. 1 engine, discharged "
   "both fire bottles, performed an air turn-back, and landed approximately 11 minutes later...\"")
kv("NTSB probable cause (ground truth):",
   "Engine fuel and control - fuel distribution - failure.")
p("Top failure clusters:")
table(["Cluster", "P(K|Q)", "Incidents"],
      [["engine fire due to component failures", "53.5%", "24"],
       ["multiple contributing mechanical failures", "9.2%", "5"],
       ["undetected fatigue cracks in engine", "7.9%", "4"],
       ["fatigue-induced engine component failures", "7.7%", "4"],
       ["nose landing gear malfunctions", "5.5%", "3"]])
p("Top causes:")
table(["Cause", "P(C|Q)"],
      [["Turbine section failure", "3.3%"],
       ["Manufacturer / production defect", "3.3%"],
       ["Turbine section wear or corrosion", "2.7%"],
       ["Inadequate inspection procedures", "2.6%"],
       ["Improper maintenance of combustion section", "2.6%"]])
p("The pipeline localizes the incident to the engine fire due to component failures cluster "
  "with 53.5% probability, a high-confidence reading of the failure area. The individual "
  "causes carry smaller probabilities (around 2-3%) because the NTSB cause taxonomy is large "
  "and the probability mass spreads across many specific codes. The cluster probability is the "
  "decision-useful level (it identifies where the problem lies), while the ranked causes are "
  "the specific items to check within that area.")

h("8.1.2 During-incident query (diagnosis + prognosis)", 3)
kv("Query (flight crew, real time):",
   "\"I just got a left-engine fire warning on takeoff. I can see flames. What could be "
   "causing this?\"")
p("Diagnosis - clusters:")
table(["Cluster", "P(K|Q)", "Incidents"],
      [["engine fire due to component failures", "53.7%", "21"],
       ["electrical failures and hazards", "9.2%", "4"],
       ["multiple contributing mechanical failures", "7.9%", "3"],
       ["undetected fatigue cracks in engine", "7.1%", "3"],
       ["bird strike during takeoff", "5.4%", "2"]])
p("Diagnosis - causes:")
table(["Cause", "P(C|Q)"],
      [["Improper maintenance installation", "3.1%"],
       ["Turbine section wear or corrosion", "3.1%"],
       ["Bird strike impact on equipment", "2.7%"],
       ["Turbine section failure", "2.5%"],
       ["Manufacturer / production defect", "2.5%"]])
p("Prognosis - most likely next-event area:")
table(["Next-event cluster", "P(K|Q)", "Sequences with a next event"],
      [["engine fire due to component failures", "55.0%", "15"],
       ["electrical failures and hazards", "9.5%", "1"],
       ["multiple contributing mechanical failures", "8.1%", "2"],
       ["undetected fatigue cracks in engine", "7.3%", "3"],
       ["bird strike during takeoff", "5.5%", "1"]])
p("From only a short crew report, diagnosis reaches essentially the same top cluster (53.7%) "
  "as the full narrative, showing the method does not require a complete investigator write-up. "
  "Prognosis then projects the most likely next-event area (engine-fire progression, 55.0%), "
  "the forward-looking information useful while the incident is still developing.")

# Fully worked example
h("Worked calculation: how the numbers are computed", 3)
p("We show the full computation for the diagnosis above, from retrieval to the final cause "
  "probability, so every number is traceable.")
p("Step 1 - Retrieve. Embed the query and retrieve the 50 most similar past incidents by "
  "similarity score.")
p("Step 2 - Group into precomputed clusters. The 50 retrieved incidents fall into their "
  "(offline) clusters. For each cluster we record how many of the 50 it holds (n) and its "
  "average similarity to the query:")
table(["Cluster (K)", "avg_sim", "n", "weight = avg_sim x n"],
      [["engine fire due to component failures", "1.013", "21", "21.28"],
       ["electrical failures and hazards", "0.916", "4", "3.67"],
       ["multiple contributing mechanical failures", "1.049", "3", "3.15"],
       ["undetected fatigue cracks in engine", "0.937", "3", "2.81"],
       ["bird strike during takeoff", "1.070", "2", "2.14"],
       ["... remaining clusters ...", "", "", "..."],
       ["Sum of all weights", "", "", "39.64"]])
p("(Note: the similarity here is the fused structural-plus-cosine score from Section 7, which "
  "is why it can slightly exceed 1.)")
p("Step 3 - Cluster probability P(K|Q) = weight / sum of weights:")
eq("P(engine fire | Q) = 21.28 / 39.64 = 0.537 (53.7%)")
eq("P(multiple mechanical | Q) = 3.15 / 39.64 = 0.079 (7.9%)")
p("Step 4 - Cause probability within a cluster, P(C|K), is a frequency count of how often a "
  "cause appears in that cluster. The cause \"turbine section failure\" appears in two clusters:")
p("   in engine fire due to component failures:  P(C|K) = 0.0286", italic=True)
p("   in multiple contributing mechanical failures:  P(C|K) = 0.1176", italic=True)
p("Step 5 - Law of total probability. The final cause probability sums across all clusters in "
  "which the cause appears:")
eq("P(C|Q) = sum over K of  P(C|K) x P(K|Q)")
eq("P(turbine section failure | Q) = (0.0286 x 0.537) + (0.1176 x 0.079)")
eq("= 0.0153 + 0.0093 = 0.0247 (2.47%)")
p("Step 6 - Rank. Repeating Step 5 for every cause and ranking produces the cause table above.")
p("Interpreting 53.7% vs 2.5%: the cluster probability (53.7%) localizes the failure area with "
  "high confidence. A specific cause is that area's confidence multiplied by the cause's share "
  "within the clusters it appears in, then summed, so it comes out smaller (2.5%). The cluster "
  "says where the problem is; the ranked causes say what to check inside it.")

# ---------- 8.2 Landing Gear ----------
h("8.2 Landing Gear", 2)

h("8.2.1 After-incident query (diagnosis only)", 3)
kv("Query (investigator narrative):",
   "\"On November 16, 2008, a deHavilland DHC-8-311 operated by Piedmont Airlines...\" "
   "(full narrative).")
kv("NTSB probable cause (ground truth):",
   "Landing gear system - gear extension and retraction - capability exceeded.")
table(["Cluster", "P(K|Q)", "Incidents"],
      [["nose landing gear malfunctions", "18.1%", "7"],
       ["multiple contributing mechanical failures", "13.2%", "7"],
       ["unexpected turbulence encounters", "13.1%", "7"],
       ["turbulence-related flight attendant injuries", "7.3%", "4"],
       ["tailwind landing misjudgments", "7.0%", "3"]])
p("Top causes:")
table(["Cause", "P(C|Q)"],
      [["Clear air turbulence effects on crew", "4.6%"],
       ["Inadvertent clear-air turbulence encounter", "3.7%"],
       ["Convective turbulence impact on crew", "3.4%"],
       ["Pilot incorrect action", "2.0%"],
       ["Bird strike impact on equipment", "2.0%"]])
p("This is a weaker result. The top cluster is only 18.1%, and the leading causes drift toward "
  "turbulence rather than landing gear. The after-incident narrative is long and mentions "
  "weather, injuries, and other events, which dilutes retrieval away from the core landing-gear "
  "failure. As shown next, the focused during-incident query corrects this. We include this "
  "case rather than omit it, to show the method's behavior is not uniformly strong and depends "
  "on query focus.")

h("8.2.2 During-incident query (diagnosis + prognosis)", 3)
kv("Query (flight crew, real time):",
   "\"I just got an unsafe nose-gear indication on approach. Three greens on the mains but "
   "nothing on the nose. What could be wrong?\"")
p("Diagnosis - clusters:")
table(["Cluster", "P(K|Q)", "Incidents"],
      [["nose landing gear malfunctions", "40.7%", "15"],
       ["landing gear structural failures", "23.9%", "9"],
       ["landing gear tire failures", "9.0%", "4"],
       ["improper pitch control during landing", "5.9%", "3"],
       ["structural integrity failures", "5.6%", "2"]])
p("Diagnosis - causes:")
table(["Cause", "P(C|Q)"],
      [["Cause not determined", "3.9%"],
       ["Nose or tail landing gear failure", "3.7%"],
       ["Landing gear actuator failure", "3.3%"],
       ["Non-optimal grain direction in part", "3.3%"],
       ["Nose landing gear fatigue failure", "3.3%"]])
p("Prognosis - next event:")
table(["Next-event cluster", "P(K|Q)", "Sequences with a next event"],
      [["nose landing gear malfunctions", "65.5%", "7"],
       ["landing gear tire failures", "14.4%", "4"],
       ["improper pitch control during landing", "9.5%", "2"],
       ["tailwind landing misjudgments", "7.0%", "1"],
       ["electrical failures and hazards", "3.5%", "1"]])
p("The focused query localizes strongly to the nose-landing-gear area (40.7%, with landing-gear "
  "structural failures next at 23.9%, together about 65% on landing gear). This is why we report "
  "all five causes rather than only the first: the top cause here is \"cause not determined,\" "
  "which is not actionable, but the next entries (nose/tail gear failure, actuator failure, "
  "fatigue failure) are the useful candidates. Showing the full ranked list lets the user skip "
  "non-informative entries.")

# ---------- 8.3 Summary ----------
h("8.3 Summary Table", 2)
table(["Example", "Query style", "Predicted top cluster", "NTSB ground truth", "Match"],
      [["Engine fire", "After-incident", "engine fire due to component failures (53.5%)",
        "engine fuel/control failure", "Correct area"],
       ["Engine fire", "During-incident", "engine fire due to component failures (53.7%)",
        "engine fuel/control failure", "Correct area"],
       ["Landing gear", "After-incident", "nose landing gear malfunctions (18.1%)",
        "landing gear extension/retraction", "Weak / mixed"],
       ["Landing gear", "During-incident", "nose landing gear malfunctions (40.7%)",
        "landing gear extension/retraction", "Correct area"]])
p("Across the focused during-incident queries, the pipeline localizes to the correct failure "
  "area. The one weak case is the long landing-gear after-incident narrative, consistent with "
  "the observation that focused, failure-centered queries retrieve more accurately than long "
  "mixed narratives.")

doc.save(OUT)
print("Saved:", OUT)
