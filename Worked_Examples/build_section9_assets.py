"""Compute Section 9 stats and generate tables (docx) + graphs (png)."""
import json
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Pt, Inches

SRC = "Worked_Examples/data/aggregate_results.jsonl"
OUTDIR = "/Users/kanushetkar/Desktop/NTSB Paper Process"

rows = [json.loads(l) for l in open(SRC) if l.strip()]
N = len(rows)

# ---- A2 vs A0 on cluster cosine ----
wins = ties = losses = 0
for r in rows:
    a0, a2 = r["sim_cluster_a0"], r["sim_cluster_a2"]
    if abs(a2 - a0) < 1e-9:
        ties += 1
    elif a2 > a0:
        wins += 1
    else:
        losses += 1
as_good = wins + ties

# ---- threshold match rates ----
def rate(key, thr):
    return sum(1 for r in rows if r[key] >= thr)

thresholds = [0.30, 0.40, 0.50]
cluster_a0 = [rate("sim_cluster_a0", t) for t in thresholds]
cluster_a2 = [rate("sim_cluster_a2", t) for t in thresholds]
cause_a0 = [rate("sim_cause_a0", t) for t in thresholds]
cause_a2 = [rate("sim_cause_a2", t) for t in thresholds]

# means
mean_cluster_a0 = np.mean([r["sim_cluster_a0"] for r in rows])
mean_cluster_a2 = np.mean([r["sim_cluster_a2"] for r in rows])
mean_cause_a0 = np.mean([r["sim_cause_a0"] for r in rows])
mean_cause_a2 = np.mean([r["sim_cause_a2"] for r in rows])

print(f"N = {N}")
print(f"A2 vs A0 cluster cosine: wins={wins}, ties={ties}, losses={losses}, as_good={as_good} ({as_good/N*100:.1f}%)")
print(f"Cluster matches A0  (>=.30/.40/.50): {cluster_a0}")
print(f"Cluster matches A2  (>=.30/.40/.50): {cluster_a2}")
print(f"Cause matches   A0  (>=.30/.40/.50): {cause_a0}")
print(f"Cause matches   A2  (>=.30/.40/.50): {cause_a2}")
print(f"Mean cluster cosine A0={mean_cluster_a0:.3f} A2={mean_cluster_a2:.3f}")
print(f"Mean cause   cosine A0={mean_cause_a0:.3f} A2={mean_cause_a2:.3f}")

# =================== GRAPHS ===================
GRAY_D = "#333333"
GRAY_L = "#999999"

# Graph 1: match rate at thresholds (cluster), A0 vs A2
fig, ax = plt.subplots(figsize=(6, 4))
x = np.arange(len(thresholds))
w = 0.35
pa0 = [c / N * 100 for c in cluster_a0]
pa2 = [c / N * 100 for c in cluster_a2]
ax.bar(x - w/2, pa0, w, label="A0 (embedding)", color=GRAY_L)
ax.bar(x + w/2, pa2, w, label="A2 (structural)", color=GRAY_D)
ax.set_xticks(x)
ax.set_xticklabels([f"cosine >= {t:.2f}" for t in thresholds])
ax.set_ylabel("% of 77 test cases")
ax.set_title("Top-cluster match rate vs NTSB probable cause")
ax.legend()
for i, v in enumerate(pa0):
    ax.text(i - w/2, v + 1, f"{v:.0f}%", ha="center", fontsize=8)
for i, v in enumerate(pa2):
    ax.text(i + w/2, v + 1, f"{v:.0f}%", ha="center", fontsize=8)
ax.set_ylim(0, 100)
plt.tight_layout()
g1 = f"{OUTDIR}/fig_cluster_match_rate.png"
plt.savefig(g1, dpi=200)
plt.close()

# Graph 2: A2 vs A0 head-to-head pie/bar
fig, ax = plt.subplots(figsize=(5, 4))
labels = ["A2 better", "Tie", "A0 better"]
vals = [wins, ties, losses]
colors = [GRAY_D, GRAY_L, "#cccccc"]
ax.bar(labels, vals, color=colors)
for i, v in enumerate(vals):
    ax.text(i, v + 0.5, str(v), ha="center")
ax.set_ylabel("Number of test cases")
ax.set_title(f"A2 vs A0 on cluster cosine (n={N})")
ax.set_ylim(0, max(vals) + 6)
plt.tight_layout()
g2 = f"{OUTDIR}/fig_a2_vs_a0.png"
plt.savefig(g2, dpi=200)
plt.close()

# Graph 3: distribution of cluster cosine similarity (A2)
fig, ax = plt.subplots(figsize=(6, 4))
sims = [r["sim_cluster_a2"] for r in rows]
ax.hist(sims, bins=15, color=GRAY_L, edgecolor=GRAY_D)
for t in thresholds:
    ax.axvline(t, color="black", linestyle="--", linewidth=1)
    ax.text(t, ax.get_ylim()[1]*0.9, f"{t:.2f}", rotation=90, fontsize=8, va="top")
ax.set_xlabel("Cosine similarity of top cluster to NTSB probable cause")
ax.set_ylabel("Number of test cases")
ax.set_title("Distribution of top-cluster similarity (A2)")
plt.tight_layout()
g3 = f"{OUTDIR}/fig_cluster_sim_hist.png"
plt.savefig(g3, dpi=200)
plt.close()

print("Saved graphs:", g1, g2, g3, sep="\n  ")

# =================== TABLES DOC ===================
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)

def table(headers, data):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, hh in enumerate(headers):
        t.rows[0].cells[i].text = hh
        for par in t.rows[0].cells[i].paragraphs:
            for run in par.runs:
                run.bold = True
    for row in data:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()

doc.add_heading("Section 9 - Tables", 1)

doc.add_heading("9.3 Aggregate results table", 2)
table(
    ["Match level (cosine threshold)", "A0 cluster", "A2 cluster", "A0 cause", "A2 cause"],
    [
        ["Solid match (>= 0.50)", f"{cluster_a0[2]} ({cluster_a0[2]/N*100:.0f}%)",
         f"{cluster_a2[2]} ({cluster_a2[2]/N*100:.0f}%)",
         f"{cause_a0[2]} ({cause_a0[2]/N*100:.0f}%)",
         f"{cause_a2[2]} ({cause_a2[2]/N*100:.0f}%)"],
        ["Plausible (>= 0.40)", f"{cluster_a0[1]} ({cluster_a0[1]/N*100:.0f}%)",
         f"{cluster_a2[1]} ({cluster_a2[1]/N*100:.0f}%)",
         f"{cause_a0[1]} ({cause_a0[1]/N*100:.0f}%)",
         f"{cause_a2[1]} ({cause_a2[1]/N*100:.0f}%)"],
        ["At least related (>= 0.30)", f"{cluster_a0[0]} ({cluster_a0[0]/N*100:.0f}%)",
         f"{cluster_a2[0]} ({cluster_a2[0]/N*100:.0f}%)",
         f"{cause_a0[0]} ({cause_a0[0]/N*100:.0f}%)",
         f"{cause_a2[0]} ({cause_a2[0]/N*100:.0f}%)"],
        ["Mean cosine", f"{mean_cluster_a0:.3f}", f"{mean_cluster_a2:.3f}",
         f"{mean_cause_a0:.3f}", f"{mean_cause_a2:.3f}"],
    ],
)
doc.add_paragraph(f"All rates are over the N = {N} held-out test incidents.")

doc.add_heading("9.4 A2 vs A0 head-to-head (cluster cosine)", 2)
table(
    ["Outcome", "Count", "% of test set"],
    [
        ["A2 closer to ground truth", wins, f"{wins/N*100:.1f}%"],
        ["Tie (identical)", ties, f"{ties/N*100:.1f}%"],
        ["A0 closer to ground truth", losses, f"{losses/N*100:.1f}%"],
        ["A2 at least as good (win+tie)", as_good, f"{as_good/N*100:.1f}%"],
    ],
)

doc.add_heading("9.5 Secondary metric (cause-level cosine)", 2)
table(
    ["Metric", "A0", "A2"],
    [
        ["Mean cause cosine", f"{mean_cause_a0:.3f}", f"{mean_cause_a2:.3f}"],
        ["Cause >= 0.40", f"{cause_a0[1]} ({cause_a0[1]/N*100:.0f}%)",
         f"{cause_a2[1]} ({cause_a2[1]/N*100:.0f}%)"],
        ["Cause >= 0.30", f"{cause_a0[0]} ({cause_a0[0]/N*100:.0f}%)",
         f"{cause_a2[0]} ({cause_a2[0]/N*100:.0f}%)"],
    ],
)

out = f"{OUTDIR}/Section 9 Tables.docx"
doc.save(out)
print("Saved tables doc:", out)
