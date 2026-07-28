"""Generate docs/jesse_jul28_update.pptx — 15-min Jesse meeting deck.

Story: start → LLM dead end → reproduce Zhang → final pipeline → held-out proof
       → structural mapping exhaustive test → why it failed → ask.

Matches tuesday_jul21 / friday_update style: short bullets, tables, bold takeaway.

Run:
  python3.11 docs/build_jesse_jul28_pptx.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

TEMPLATE = Path("/Users/kanushetkar/Downloads/Copy of NTSB project update 6_30_26.pptx")
HERE = Path(__file__).resolve().parent
OUT = HERE / "jesse_jul28_update.pptx"
ARCH_PNG = HERE / "figures" / "narrative_to_bn_architecture.png"

prs = Presentation(str(TEMPLATE)) if TEMPLATE.is_file() else Presentation()
if TEMPLATE.is_file():
    xml_slides = prs.slides._sldIdLst
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for sld in list(xml_slides):
        prs.part.drop_rel(sld.get(R))
        xml_slides.remove(sld)

LAYOUTS = {l.name: l for l in prs.slide_layouts}
TITLE_LAYOUT = LAYOUTS.get("TITLE", prs.slide_layouts[0])
OBJECT_LAYOUT = LAYOUTS.get("OBJECT", prs.slide_layouts[1])


def _placeholder(slide, want_title):
    for ph in slide.placeholders:
        is_title = ph.placeholder_format.idx == 0 or "title" in ph.name.lower()
        if is_title == want_title:
            return ph
    return None


def add_content_slide(title, bullets=None, note=None):
    s = prs.slides.add_slide(OBJECT_LAYOUT)
    t = _placeholder(s, True)
    if t is not None:
        t.text_frame.text = title
    body = _placeholder(s, False)
    if bullets and body is not None:
        tf = body.text_frame
        tf.clear()
        first = True
        for item in bullets:
            text, level = item[0], item[1]
            opts = item[2] if len(item) > 2 else {}
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.level = level
            r = p.add_run()
            r.text = text
            if opts.get("bold"):
                r.font.bold = True
    elif body is not None:
        body.text_frame.clear()
    if note:
        s.notes_slide.notes_text_frame.text = note
    return s


def add_table(slide, rows, left, top, width, height, col_widths=None, font_size=13):
    nrows, ncols = len(rows), len(rows[0])
    gfx = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    tbl = gfx.table
    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = Inches(w)
    for r in range(nrows):
        for c in range(ncols):
            cell = tbl.cell(r, c)
            para = cell.text_frame.paragraphs[0]
            run = para.add_run()
            run.text = str(rows[r][c])
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if r == 0:
                run.font.bold = True
    return tbl


def add_takeaway(slide, text, top=5.35, size=16):
    tb = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True


# ---- 1 TITLE
s = prs.slides.add_slide(TITLE_LAYOUT)
tp = _placeholder(s, True)
if tp is not None:
    tp.text_frame.text = "NTSB Narrative Project — Full Story"
sp = _placeholder(s, False)
if sp is not None:
    sp.text_frame.text = "Kanu Shetkar  •  Jesse Spencer-Smith  •  July 28, 2026"

# ---- 2 MAP
add_content_slide(
    "Today's map (15 min)",
    [
        ("Where we started: Zhang's network + your structural-mapping idea", 0),
        ("Dead end #1: LLM computes probabilities directly", 0),
        ("What we built: LLM reads → frozen network calculates", 0),
        ("Proof: reproduce Zhang + 296 held-out accidents (~93% injury)", 0),
        ("Your idea tested exhaustively → appendix negative result", 0),
    ],
    note="One arc: reproduce calculator, narrative bridge, fair test of struct mapping.")

# ---- 3 ONE PICTURE
s = add_content_slide(
    "The one picture: build time vs query time",
    [
        ("BUILD (once): coded data → frozen BN; narratives → embedding index", 0),
        ("QUERY: counting path (one condition) OR network path (many facts)", 0),
        ("Rule: narrative never retrains the network", 0),
    ],
    note="Undergrad analogy: calculator trained on old accidents, frozen; "
         "new stories only fill in the form.")
if ARCH_PNG.is_file():
    s.shapes.add_picture(str(ARCH_PNG), Inches(7.0), Inches(1.55), height=Inches(3.6))
add_takeaway(s, "The LLM is the reader. The network is the calculator.")

# ---- 4 LLM DEAD END
add_content_slide(
    "Dead end #1: LLM-only probabilities",
    [
        ("Asked GPT for probabilities directly → ~10x wrong vs Zhang", 0),
        ("One word change → different answer (not reproducible)", 0),
        ("Lesson: LLM = translator; network = calculator", 0),
    ],
    note="This motivates Section 3.5 / tiered parser design.")

# ---- 5 REPRODUCTION
s = add_content_slide(
    "Reproduction: what we fixed before trusting narratives",
    note="Table 7 85/85; scoreboard 77/93 after upgrades.")
add_table(s, [
    ["Issue", "Fix", "Result"],
    ["Fire count 38→102", "Merge occurrence files", "Denom 102 ✓"],
    ["Table 7 partial", "Cause/Factor filter", "85/85 exact"],
    ["No pilot nodes", "Person findings", "Fig 12 runs"],
    ["Boolean injury", "4-level severity", "Held-out eval"],
], left=0.7, top=1.5, width=11.8, height=3.2,
   col_widths=[3.2, 4.0, 4.6], font_size=14)
add_takeaway(s, "Trust the calculator first, then test narratives.", top=5.0)

# ---- 6 PIPELINE
add_content_slide(
    "Final pipeline: tiered parser → evidence → frozen BN",
    [
        ("Tier 1: NTSB vocabulary match (11/11 scenarios; 10/11 no LLM)", 0),
        ("Tier 2: LLM only when Tier 1 empty; strengths from data", 0),
        ("Hard + soft + stated evidence → propagation", 0),
    ],
    note="outputs/combined_parser_validation.md")

# ---- 7 HELD-OUT
s = add_content_slide(
    "Does it work? 296 held-out accidents (2007–2019)",
    note="Never used in network construction.")
add_table(s, [
    ["Method", "Injury", "Damage"],
    ["Prior (no narrative)", "58%", "43%"],
    ["Full narrative → BN", "89%", "64%"],
    ["Best readout (narr-sev)", "93%", "81%"],
    ["Supervised LR baseline", "88%", "65%"],
    ["LLM-only front door", "69%", "58%"],
], left=1.0, top=1.55, width=11.0, height=3.0,
   col_widths=[5.0, 2.5, 2.5], font_size=15)
add_takeaway(s, "Narratives early; coded labels late — ~93% injury on unseen cases.", top=4.85)

# ---- 8 JESSE IDEA
add_content_slide(
    "Your idea: structural mapping",
    [
        ("LLM extracts causal chain (role / system / mechanism steps)", 0),
        ("Align chains (Needleman-Wunsch); rerank by structure + cosine", 0),
        ("Goal: find 'same mechanism, different words' cases cosine misses", 0),
    ],
    note="Frame as reasonable hypothesis we tested to falsify properly.")

# ---- 9 WHAT WE TESTED
s = add_content_slide(
    "Structural mapping: every variant we tested",
    note="outputs/structmap_final_verdict/REPORT.md")
add_table(s, [
    ["Variant", "What it tests"],
    ["A0 embedding baseline", "Cosine similarity only"],
    ["A2 rerank top-50", "Your original rerank idea"],
    ["V1 best-match bypass", "Drop cluster averaging"],
    ["V3 wider pool rerank", "Change pool then rerank"],
    ["V4 struct-first ALL 1,703", "Retrieve by structure alone"],
    ["V5 hybrid embedding", "Narrative ⊕ chain in one vector"],
    ["V6 struct severity vote", "Structure on 296 held-out task"],
], left=0.6, top=1.45, width=12.0, height=3.8,
   col_widths=[3.5, 8.5], font_size=13)
add_takeaway(s, "Rerank, replace retrieval, hybrid — every insertion point.", top=5.55)

# ---- 10 RESULTS
s = add_content_slide(
    "Structural mapping: results",
    note="Never beat baseline; struct-first hurts severity.")
add_table(s, [
    ["Test", "Result"],
    ["Rerank A2 vs A0 (diagnosis)", "No significant win"],
    ["V1 best-match bypass", "Harmful −1.6 to −2.6 pp (p<0.001)"],
    ["V4 struct-first (severity)", "Injury 92.9%→81.1%; damage 81.4%→70.9%"],
    ["Best struct cell (1 of 28)", "+1.4 pp damage, p=0.125 (noise)"],
    ["k=50 no structure control", "+1.0 pp damage (same noise)"],
], left=0.8, top=1.5, width=11.5, height=3.3,
   col_widths=[4.5, 7.0], font_size=14)
add_takeaway(s, "Exhaustively null. Best struct win = pool-size jitter.", top=5.1)

# ---- 11 WHY
add_content_slide(
    "Why structural mapping did not help",
    [
        ("Chains too generic (~5 steps) — many accidents share the same shape", 0),
        ("Embedding already encodes most signal (struct ρ ≈ 0.01 vs truth)", 0),
        ("When struct disagrees with embedding, it is wrong more than right", 0),
        ("Averaging 50 neighbors beats picking one 'best struct' incident", 0),
    ],
    note="Honest to Jesse: idea testable; this extraction stack is information-poor.")

# ---- 12 ASK
add_content_slide(
    "Where we are + question for you",
    [
        ("Paper: reproduce Zhang + narrative bridge + held-out 296 + ablations", 0),
        ("Struct mapping → Appendix A (clean negative result)", 0),
        ("", 0),
        ("OK to freeze struct mapping as closed?", 0, {"bold": True}),
        ("Anything you want rerun before we finalize the draft?", 0, {"bold": True}),
    ],
    note="Stop and let Jesse direct.")

prs.save(str(OUT))
print(f"Wrote {OUT}")
