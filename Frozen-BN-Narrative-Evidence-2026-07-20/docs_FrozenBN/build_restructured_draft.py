#!/usr/bin/env python3
"""Build the RESTRUCTURED paper draft.

Takes the user's current draft (Draft NTSB Paper 7/21/26.docx), physically
moves every surviving paragraph/table/figure into the new Zhang-style
section order, and inserts yellow-highlighted GUIDE blocks above each
section explaining: what belongs there, what is already present, what to
add, what to change, and which references to use (cross-referenced to
docs/paper_revision_checklist.docx item numbers).

Output: docs/Draft NTSB Paper RESTRUCTURED 7_22_26.docx  (copy to Desktop after)

Run:
  /Library/Frameworks/Python.framework/Versions/3.11/bin/python3.11 \
      docs/build_restructured_draft.py
"""
from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt

SRC = "/Users/kanushetkar/Desktop/ Draft NTSB Paper 7:21:26.docx"
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Draft NTSB Paper RESTRUCTURED 7_22_26.docx"

doc = docx.Document(SRC)
body = doc.element.body
children = list(body.iterchildren())
paras = [c for c in children if c.tag == qn("w:p")]
sect_pr = body.find(qn("w:sectPr"))
pos = {id(el): i for i, el in enumerate(children)}


def block(a: int, b: int) -> list:
    """All body elements from paragraph a to paragraph b inclusive
    (tables and figure paragraphs in between come along)."""
    i0, i1 = pos[id(paras[a])], pos[id(paras[b])]
    return [el for el in children[i0:i1 + 1] if el is not sect_pr]


def heading(text: str, level: int = 1):
    try:
        p = doc.add_paragraph(text, style=f"Heading {level}")
    except KeyError:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(16 if level == 1 else 13)
    return [p._p]


def guide(lines: list[str]) -> list:
    """Yellow-highlighted instruction block. First line bold."""
    els = []
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if i == 0:
            r.bold = True
        els.append(p._p)
    return els


order: list = []

# ----------------------------------------------------------------- top note
order += guide([
    "[HOW TO USE THIS DOCUMENT - delete this block when done]",
    "This is your July 21 draft restructured into the new section order, "
    "which mirrors Zhang & Mahadevan's RESS paper section for section. "
    "Every yellow block is an instruction: it says what the section is for, "
    "what text you already have (moved here for you), what you still need "
    "to write, what to change, and which references go where. Delete each "
    "yellow block as you finish that section.",
    "\"Item N\" and reference details refer to docs/paper_revision_checklist.docx "
    "(open it side by side - it has every link and how to verify each source).",
    "Section numbers in NEW headings are correct. Some paragraphs moved from "
    "old sections still carry old subsection numbers (e.g. 6.2, 5.1) - "
    "renumber those as the guides tell you.",
    "Writing rule from Maha: never paste AI prose. Read the guide, write the "
    "text in your own words, then send it back for fact-check.",
])

# --------------------------------------------------------------- title block
order += block(0, 3)
order += guide([
    "[TITLE BLOCK]",
    "CHANGE: update the date line. Decide the final title with Maha - the "
    "working title still describes the old retrieval-only project. A title "
    "that covers the real contribution would mention narratives driving a "
    "Bayesian network, e.g. something like: 'Driving a Bayesian network "
    "for aviation risk assessment directly from unstructured accident "
    "narratives'. Put it in your own words.",
])

# ------------------------------------------------------------------ abstract
order += guide([
    "[ABSTRACT - REWRITE LAST, after the whole body is done]",
    "WHAT: one paragraph, unnumbered, before Section 1.",
    "HAVE: the old abstract below - it only describes the counting path "
    "(embeddings, clusters, diagnosis/prognosis), wrong dates, and 'two "
    "similar methodologies' (Items 1-4).",
    "ADD when you rewrite: (1) we reproduce Zhang's Bayesian network from "
    "the coded data; (2) we add a narrative-to-evidence bridge (two-tier "
    "parser: deterministic first, LLM fallback with data-measured "
    "strengths); (3) held-out validation on 296 unseen accidents "
    "(2007-2019) where the narrative-driven network matches or beats a "
    "supervised baseline.",
    "CHANGE: 1982-2019 everywhere (Item 1); causes are UPSTREAM, next "
    "events are DOWNSTREAM (Item 3); fix the Streamlit phrase (Item 4).",
])
order += block(4, 6)

# ------------------------------------------------------------ 1 Introduction
order += heading("1. Introduction", 1)
order += guide([
    "[SECTION 1 GUIDE - Introduction]",
    "WHAT: motivation, the NTSB data opportunity, short literature review, "
    "then a numbered contribution list and a roadmap paragraph. Zhang has "
    "no separate Related Work section - his review lives inside the intro, "
    "so your old Section 2.1 (Zhang description) has been moved in below.",
    "MOVED OUT for you: the fault-tree/event-tree and Bayesian-network "
    "paragraphs now live in Section 3 (Background); the LLM-overconfidence "
    "paragraphs now live in Section 3.5 (design rationale). Add ONE "
    "sentence here that traditional tools (fault trees, event trees, BNs) "
    "are reviewed in Section 3.",
    "FIX in the paragraphs below: Federal Aviation ADMINISTRATION (Item 5); "
    "replace '(Ref)' with the FAA Air Traffic by the Numbers page (Item 6); "
    "replace '(re)' with the IATA Annual Safety Report (Item 7); 'The NTSB "
    "consisted of data' -> the NTSB DATABASE contains, mission IS (Item 8); "
    "2,243 ACCIDENTS 1982-2019 + cite the FAA Part 121 page (Item 9); cite "
    "Zhang right where you say earlier BN analyses used coded fields "
    "(Item 10); cite an NLP-for-aviation paper for the unstructured-text "
    "opportunity (Item 11).",
    "ADD (new writing): (a) one breadth paragraph - BNs in transportation "
    "safety (Ale et al. CATS, Luxhoj 2003, Ancel et al. 2015 - checklist "
    "refs 13-15) and ML for aviation risk (Zhang & Mahadevan DSS 2019 - "
    "ref 12 - which pairs nicely with your supervised baseline); (b) a "
    "numbered contribution list (reproduction of Zhang's BN; the narrative "
    "bridge; held-out validation); (c) a roadmap paragraph 'Section 2 "
    "describes... Section 3... etc.' exactly like Zhang ends his intro "
    "(Item 15).",
    "The Zhang paragraphs below came from old Section 2.1: fix 'prior "
    "possible probabilities' (Item 17) and delete the sentence duplicated "
    "from the abstract (Item 16).",
])
order += block(8, 10)
order += block(22, 25)

# --------------------------------------------------- 2 NTSB data
order += heading("2. NTSB accident investigation data", 1)
order += guide([
    "[SECTION 2 GUIDE - Data. Mirrors Zhang's Section 2: data comes "
    "BEFORE methods.]",
    "HAVE: your old Section 3 paragraphs, plus the Merged Dataset "
    "paragraph moved in from old 6.1 (delete its '6.1 Merged Dataset' "
    "heading - it is just a paragraph here now).",
    "FIX: 'sequence of events... that had recovered' -> 'that were "
    "RECORDED' (Item 24); accidents not incidents (G2); 1982-2019 (G3); "
    "pick ONE consistent file-count story - the draft says 14 files in "
    "one place and eight in another (Item 23).",
    "ADD (new writing, most important thing in this section): the split "
    "definition (Item 25). State plainly: 2,243 accidents 1982-2019; the "
    "network is built ONLY from the 1,742 accidents in the 1982-2006 "
    "window (matching Zhang); 296 accidents from 2007-2019 with usable "
    "factual narratives are held out and never touched during building - "
    "they are the test set in Section 5.3. Also add 2-3 sentences "
    "distinguishing CODED records (occurrence/finding codes, take months "
    "to produce) from NARRATIVES (free text, available immediately) - "
    "this distinction is the whole premise of the paper.",
    "REFERENCES: NTSB database download page data.ntsb.gov/avdata; "
    "Zhang's GitHub repository as a web reference with access date "
    "(checklist ref 22 - verify it is still public); NTSB Aviation Coding "
    "Manual codman.pdf (checklist ref 23) where the code dictionary is "
    "described.",
])
order += block(30, 31)
order += block(61, 62)

# --------------------------------------------------- 3 Background
order += heading("3. Background: Bayesian networks, embeddings, and large "
                 "language models", 1)
order += guide([
    "[SECTION 3 GUIDE - Background. Mirrors Zhang's Section 3 (his is "
    "titled 'Bayesian networks'); yours also needs embeddings + LLMs "
    "because your method uses them.]",
    "HAVE (moved in for you): the fault-tree/event-tree paragraph and the "
    "Bayesian-network paragraph from your old intro, then the embedding / "
    "k-means / law-of-total-probability subsections from old Section 4. "
    "Renumber the old '4.x' subsection headings to 3.1, 3.2, 3.3...",
    "FIX in fault/event tree paragraph: 'contradictory' -> COMPLEMENTARY "
    "(fault trees reason backward to causes, event trees forward to "
    "outcomes) and cite the NASA PRA Procedures Guide at the end of the "
    "definitions (Item 12).",
    "FIX in BN paragraph(s): the definition appears twice in the draft - "
    "merge into one paragraph and cite Pearl 1988 + Jensen 1996 "
    "(Item 13).",
    "FIX in embeddings subsection: 'near1' spacing, 'multi-dimensional "
    "mode' -> SPACE, cosine range statement, dot product in the equation, "
    "cite the OpenAI embeddings docs (Items 27-28). Put ALL equations in "
    "the Word equation editor and number them (1), (2), ... - every "
    "equation in the draft lost its Greek symbols (G4).",
    "FIX in k-means subsection: delete the IBM marketing URL; cite "
    "MacQueen 1967 (original k-means) and scikit-learn (implementation) "
    "(Item 29).",
    "MOVED OUT: Needleman-Wunsch is now in Appendix A - it is only used "
    "by structural mapping, which is no longer a main-line method.",
])
order += block(11, 12)
order += block(33, 41)
order += heading("3.5 Design rationale: why not an LLM-only pipeline?", 2)
order += guide([
    "[SECTION 3.5 GUIDE - THIS IS THE PART MAHA EXPLICITLY ASKED FOR: the "
    "full-LLM attempt and why it did not work, placed BEFORE the "
    "methodology so the reader knows why the design is tiered.]",
    "HAVE (moved in for you): the LLM-overconfidence paragraphs from your "
    "old intro, plus your old Section 5 (Initial LLM Only Approach: "
    "pipeline / limitations / lesson).",
    "CHANGE: compress old Section 5 from three subsections into about "
    "three paragraphs of prose - delete the '5.1/5.2/5.3' headings. "
    "Formalize the tone: 'LLMs' probabilities' or 'LLM-generated "
    "probabilities' (grammar), no conversational phrasing like 'It's too "
    "confident' (Item 14). Cite Ji et al. 2023 (ACM Computing Surveys) "
    "for the hallucination claim (checklist ref 11).",
    "END the subsection with the two design principles that the failures "
    "taught (your words): (1) a deterministic parser runs first and the "
    "LLM only fills gaps; (2) every LLM-proposed fact gets its strength "
    "measured from historical data, never from the LLM's own confidence. "
    "Say the quantitative comparison appears in Section 5.5.",
    "Keep the story QUALITATIVE here - the numbers (LLM-first vs tiered "
    "on 296 accidents) belong in Results 5.5.",
])
order += block(13, 18)
order += block(45, 59)

# --------------------------------------------------- 4 Methodology
order += heading("4. Proposed methodology", 1)
order += guide([
    "[SECTION 4 GUIDE - mirrors Zhang's Section 4. Subsection 4.2 keeps "
    "your old Section 6 text; 4.1, 4.3, 4.4, 4.5 are NEW WRITING - this "
    "is the biggest writing job in the paper. Insert the two-lane "
    "architecture figure (docs/figures/narrative_to_bn_architecture.png) "
    "at the START of this section and walk the reader through it: build "
    "time on top, query time below.]",
])
order += heading("4.1 Reproducing Zhang's Bayesian network construction", 2)
order += guide([
    "[SECTION 4.1 GUIDE - ALL NEW. Source material: tests/bn_upgraded.py "
    "and docs/ALL_TABLES_EXACT_COMPARISON.md.]",
    "WRITE: (a) Zhang's recipe in your own words - event-sequence graphs "
    "per accident, prior probabilities from accident counts over BTS "
    "departure data (cite the Bureau of Transportation Statistics, "
    "checklist ref 30), conditional probabilities via his Beta-CDF "
    "function, assembly and pruning into one network; (b) your "
    "implementation in pyAgrum (cite Ducamp et al. 2020, checklist ref "
    "28); (c) your two upgrades and WHY: person-finding nodes (pilot / "
    "crew causes Zhang's coding drops), and multi-state severity nodes "
    "(injury: fatal/serious/minor/none; damage: destroyed/substantial/"
    "minor/none) so the network predicts outcomes at the granularity the "
    "held-out evaluation needs.",
    "This subsection is what makes the reproduction claim in Results 5.1 "
    "credible - be precise, cite Zhang's equation numbers where you "
    "follow them.",
])
order += heading("4.2 Narrative retrieval and the counting path", 2)
order += guide([
    "[SECTION 4.2 GUIDE - your old Section 6 text, kept below.]",
    "CHANGE: renumber the old '6.2 Embeddings ... 6.6 Prognosis' headings "
    "to 4.2.1 ... 4.2.5. (Old 6.1 Merged Dataset moved to Section 2.) "
    "Rebuild every formula in the equation editor with real summation "
    "symbols and number them (G4). Check every number (cluster count, "
    "top-50, dataset sizes) against the refined dataset - some were "
    "computed on the old data.",
    "ADD one framing sentence at the top: this is the COUNTING PATH - it "
    "answers one condition at a time by weighted counting over similar "
    "accidents (law of total probability); the NETWORK PATH (4.3-4.4) "
    "handles several facts at once through the Bayesian network.",
    "DELETE any reference to the fused structural score (that belonged "
    "to the old Section 7, now Appendix A).",
])
order += block(63, 80)
order += heading("4.3 The narrative-to-evidence parser (two tiers)", 2)
order += guide([
    "[SECTION 4.3 GUIDE - ALL NEW. Source material: llm_evidence.py and "
    "outputs/combined_parser_validation.md.]",
    "WRITE: (a) Tier 1, the deterministic parser - exact matching against "
    "the official NTSB coded vocabulary (cite the Aviation Coding Manual, "
    "checklist ref 23); facts it finds enter as HARD evidence with "
    "confidence 1.0; when the vocabulary is present this tier alone "
    "reproduces clicked evidence exactly. (b) Tier 2, the LLM tier "
    "(gpt-4.1) - fires ONLY when Tier 1 finds no event facts; it maps "
    "described facts to vocabulary nodes (e.g. 'the first officer's "
    "windshield cracked' -> the window node). (c) The grounding rule: "
    "every LLM-proposed fact gets its strength measured from the ~100 "
    "most similar historical accidents - how often that fact truly "
    "appears in accidents like this one - NEVER from the LLM's own "
    "confidence. This is the direct answer to the Section 3.5 failures.",
])
order += heading("4.4 How evidence enters the network", 2)
order += guide([
    "[SECTION 4.4 GUIDE - ALL NEW. Source material: the architecture "
    "figure and llm_evidence.py.]",
    "WRITE the three evidence types and their mechanics: (a) HARD "
    "evidence - fact named in vocabulary, node set directly to the "
    "observed state; (b) SOFT evidence - described facts enter via "
    "virtual evidence / Jeffrey conditioning at the measured strength "
    "(cite Jeffrey, The Logic of Decision, checklist ref 27 - and Pearl "
    "again for virtual evidence); (c) STATED SEVERITY - outcomes the "
    "narrative mentions directly, weighted by how reliable stated "
    "outcomes historically are. Then one paragraph: with evidence set, "
    "the network computes updated probabilities for causes and for "
    "injury/damage severity. (Terminology warning from the meeting: only "
    "use 'posterior' where you mean the probability AFTER evidence is "
    "entered, and be ready to defend the word - otherwise say 'updated "
    "probability'.)",
])
order += heading("4.5 Summary", 2)
order += guide([
    "[SECTION 4.5 GUIDE - NEW, short. Zhang ends his methodology with a "
    "numbered summary (his 4.5); mirror it.]",
    "WRITE one numbered list retracing the pipeline: (1) build the "
    "network from 1982-2006 coded data per Zhang + two upgrades; (2) "
    "embed and cluster all narratives; (3) at query time, parse the "
    "narrative with the two-tier parser; (4) enter hard/soft/stated "
    "evidence; (5) read out cause and severity probabilities; counting "
    "path available for single-condition questions.",
])

# --------------------------------------------------- 5 Results
order += heading("5. Computational results", 1)
order += guide([
    "[SECTION 5 GUIDE - mirrors Zhang's Section 5. Subsections 5.1-5.4 "
    "are NEW WRITING from existing result files; 5.5 keeps your old "
    "Section 9 text (compressed).]",
])
order += heading("5.1 Reproduction results", 2)
order += guide([
    "[SECTION 5.1 GUIDE - ALL NEW. Sources: docs/TABLE7_FULL_REPRODUCTION"
    ".md, docs/ALL_TABLES_EXACT_COMPARISON.md, outputs/BN_UPGRADED_"
    "ENVELOPE.md.]",
    "WRITE: counting layer - Table 7 reproduced 85/85 exactly. Network "
    "layer - of Zhang's 93 published probabilities: 48 exact, 29 close, "
    "12 differ, 4 qualitative-only; the 12 differences sit inside the "
    "30-seed variance envelope of his own construction; on the 4 "
    "queryable cells your build is as close or closer to his published "
    "values than his own released model file. One summary table of this "
    "scoreboard. This section plays the role of Zhang's 'parameter "
    "calibration' - it proves the machinery is right before the new "
    "claims.",
])
order += heading("5.2 Parser validation", 2)
order += guide([
    "[SECTION 5.2 GUIDE - ALL NEW. Source: outputs/combined_parser_"
    "validation.md.]",
    "WRITE the three checks: identity (the 11 scoreboard sentences "
    "reproduce clicked-evidence results 11/11, Tier 1 only - the LLM "
    "never fires when vocabulary is present); safety (6 adversarial "
    "probes - unrelated/nonsense text adds no false evidence); "
    "paraphrase recovery (Tier 2 recovers described facts that Tier 1 "
    "misses, vs the retrieval-only fallback). Small table of the three "
    "checks.",
])
order += heading("5.3 Held-out evaluation on 296 unseen accidents", 2)
order += guide([
    "[SECTION 5.3 GUIDE - ALL NEW. THE HEADLINE TABLE. Sources: outputs/"
    "heldout_significance.md and outputs/heldout_narrative_bn_eval.json.]",
    "WRITE: setup (296 accidents 2007-2019, never used in building; "
    "predict injury severity fatal/serious/minor/none and damage "
    "destroyed/substantial/minor/none from the narrative alone). The "
    "table: every predictor (network prior alone, hard only, soft+stated, "
    "full, narrative-severity readout, LLM-tier, supervised logistic-"
    "regression baseline) with accuracy and 95% bootstrap CIs; McNemar "
    "paired tests for the key comparisons.",
    "REFERENCES here: Brier 1950 for the Brier score (checklist ref 31), "
    "McNemar 1947 for the paired test (ref 32), scikit-learn for the LR "
    "baseline (ref 26).",
    "BE HONEST in the prose (Maha will check): damage via the BN chain "
    "ties, not beats, the LR baseline; the LLM tier is below the "
    "deterministic path on injury. The claim that survives scrutiny: the "
    "narrative-driven network MATCHES OR BEATS a supervised model that "
    "needed labeled training data - while ours needs none at query time.",
])
order += heading("5.4 Scenario analysis: one real held-out accident", 2)
order += guide([
    "[SECTION 5.4 GUIDE - ALL NEW. Source: outputs/real_narrative_"
    "walkthrough.md. Named 'scenario analysis' deliberately to echo "
    "Zhang's Section 5.3.]",
    "WRITE one accident end to end: the narrative text, which tier fired "
    "and why, the evidence list with strengths, the network's "
    "probabilities, and what actually happened. If space allows add the "
    "second walkthrough (one Tier-1 case, one Tier-2 case). This replaces "
    "the old Section 8 worked examples (now Appendix B).",
])
order += heading("5.5 Ablations: LLM-only front door and structural "
                 "reranking", 2)
order += guide([
    "[SECTION 5.5 GUIDE - two parts: (a) NEW LLM-first comparison, (b) "
    "your old Section 9 compressed. The old Section 9 text is kept below "
    "as raw material.]",
    "PART (a) - NEW, write first (this is the evidence for Maha's "
    "requested Section 3.5): LLM-first (LLM as the ONLY reader, no "
    "deterministic tier) vs the tiered design on the same 296 accidents. "
    "Report: comparable top-1 accuracy but worse calibration (higher "
    "log-loss); failed exact reproduction on 1 of the 11 identity "
    "scenarios (a fact that should enter hard came in soft); "
    "run-to-run nondeterminism; an API call on every query. One small "
    "table, back-reference: 'consistent with the failures that motivated "
    "the design (Section 3.5)'. Source: outputs/heldout_significance.md "
    "(llm-first rows) and outputs/llm_only_identity.md.",
    "PART (b) - COMPRESS the old Section 9 below into ONE paragraph + "
    "pointer to Appendix A: early in the project we tested whether "
    "structure-aware reranking (LLM-extracted causal chains, sequence "
    "alignment, score fusion) improves cause retrieval; on the original "
    "evaluation top-1 went 32.5% -> 33.8% (25 vs 26 of 77), McNemar "
    "p = 1.000 - not significant; takeaway: embedding similarity already "
    "carries most of the structural signal, so the extra LLM cost is not "
    "justified and it was dropped from the final pipeline.",
    "PLACEHOLDER: a rerun on the full refined corpus with the same "
    "296-accident split is in progress - swap in the new numbers when "
    "ready and change '77 cases' to the new count. Mark it [UPDATE].",
    "Delete the old 9.1-9.5 subsection headings; figures 9.1/9.2 move to "
    "Appendix A or get cut. Fix 'casual' -> 'causal' everywhere "
    "(Item 18).",
])
order += block(181, 206)

# --------------------------------------------------- 6 Discussion
order += heading("6. Discussion", 1)
order += guide([
    "[SECTION 6 GUIDE - restructure to Zhang's exact scheme: 6.1 "
    "Contributions, 6.2 Limitations, 6.3 Future work. NO separate "
    "conclusion - Zhang does not have one; fold your old '10.1 "
    "Conclusion' into 6.1. Your old Section 10 text is below as raw "
    "material.]",
    "6.1 Contributions MUST OPEN with the 'what did we gain' answer - "
    "the question Maha pressed hardest in the meeting: coded NTSB "
    "records take months to be produced after an accident; narratives "
    "exist from day one; this pipeline lets Zhang's network run on "
    "day-one text with accuracy that matches or beats a supervised "
    "baseline trained on labeled data. Then list: validated reproduction "
    "(+ documented differences), the grounded two-tier parser, held-out "
    "validation with significance tests.",
    "6.2 Limitations - keep the honesty, update the content: the old "
    "text says 'test set was small' (was 77; now 296 - update); add: "
    "damage prediction ties the LR baseline; LLM tier below the "
    "deterministic path on injury; single database (NTSB Part 121), "
    "English narratives only; LLM tier needs an API at query time.",
    "6.3 Future work: structural mapping at scale (cite the rerun once "
    "done), richer evidence types, other transportation databases, "
    "prognosis through the network.",
    "DELETE from the old text below anything that undercuts the paper "
    "('we get close - the general area' as the summary of the whole "
    "project) - that described the OLD pipeline; the new results say "
    "more, so let them speak.",
])
order += block(208, 216)

# --------------------------------------------------- Declarations
order += heading("Declarations", 1)
order += guide([
    "[DECLARATIONS GUIDE - ALL NEW, required by the journal. Zhang's "
    "paper has CRediT, competing interests, and acknowledgment blocks - "
    "mirror them.]",
    "WRITE four short blocks: (1) CRediT author statement (draft with "
    "Maha: Kanu - methodology, software, validation, writing; Mahadevan "
    "- conceptualization, supervision, review; Spencer-Smith - "
    "supervision, review). (2) Declaration of competing interest "
    "(standard sentence: none). (3) Data availability: NTSB data are "
    "public; state whether your code will be shared on request or via "
    "repository - decide with Maha. (4) Declaration of generative-AI "
    "use, per Elsevier policy - the wording we agreed on: AI tools were "
    "used to assist with software development and editing suggestions; "
    "all methodology, analysis, and final text are the authors'; the "
    "authors reviewed and take full responsibility for all content. Put "
    "in your own words / journal's required format.",
])

# --------------------------------------------------- References
order += heading("References", 1)
order += guide([
    "[REFERENCES GUIDE]",
    "HAVE: 5 entries below - the paper needs roughly 30. The complete "
    "list with links, where each goes, and how to verify each one is in "
    "docs/paper_revision_checklist.docx (refs 1-32). Work through it as "
    "you write each section.",
    "RULES: numbered in order of FIRST appearance in the text - do the "
    "final renumbering pass at the very END of drafting, not before. Web "
    "references need: Title. Year. URL. Accessed [date]. Papers/books "
    "need: authors, title, journal/publisher, year, volume/pages - no "
    "bare links.",
    "OPEN ITEM: the 'Spencer-Smith and Goldstone, 1997' citation could "
    "not be verified - ASK JESSE directly for the exact reference before "
    "submitting (Item 20).",
])
order += block(218, len(paras) - 1)

# --------------------------------------------------- Appendix A
order += heading("Appendix A. Structural mapping: method details "
                 "(earlier exploration)", 1)
order += guide([
    "[APPENDIX A GUIDE - everything structural-mapping lives here now: "
    "the Gentner motivation (old 2.2), Needleman-Wunsch background (old "
    "4.4), the method (old Section 7), and any Section 9 detail you cut "
    "from 5.5.]",
    "WHY an appendix: the main text needs only the one-paragraph summary "
    "in 5.5; a null result does not earn 5 pages of main-text real "
    "estate, but the appendix preserves the work for a curious reviewer.",
    "IF the page budget is tight, deleting this whole appendix is fine - "
    "the 5.5 paragraph stands alone.",
    "FIX before keeping: 'casual chains' -> CAUSAL everywhere (Item 18); "
    "'we retrieved and reranked our embeddings from Nogueira and Cho' -> "
    "you adopted their RERANKING idea, not their embeddings (Item 19); "
    "rebuild formulas 1-4 in the equation editor - the alpha and "
    "summation symbols are missing (G4); renumber old 7.x headings to "
    "A.1, A.2, ...",
    "CITATIONS for this appendix (all in the checklist): Gentner 1983, "
    "Falkenhainer 1989, Goldstone 1994, Needleman & Wunsch 1970, "
    "Nogueira & Cho 2019, and the unresolved Spencer-Smith 1997 (ask "
    "Jesse).",
])
order += block(27, 28)
order += block(42, 43)
order += block(82, 113)

# --------------------------------------------------- Appendix B
order += heading("Appendix B. Worked examples from the retrieval pipeline", 1)
order += guide([
    "[APPENDIX B GUIDE - your old Section 8 (engine fire + landing gear "
    "examples), kept intact below.]",
    "CAUTION - these numbers were computed with the FUSED structural+"
    "cosine score (the worked calculation even says so in Step 2), which "
    "the final pipeline does not use. Two options: (a) recompute the two "
    "examples with the final cosine-only pipeline and update every "
    "number, or (b) delete this appendix entirely and rely on the new "
    "5.4 walkthrough. Option (b) is less work and loses little - decide "
    "with Maha. Do NOT leave it as is: a reviewer who recomputes Step 2 "
    "would get different numbers.",
    "If you keep it: renumber 8.x headings to B.1, B.2...; fix equations "
    "(G4); 'incidents' -> accidents where applicable (G2).",
])
order += block(115, 179)

# --------------------------------------------------- rebuild body
for child in list(body.iterchildren()):
    body.remove(child)
seen = set()
for el in order:
    if id(el) in seen:
        continue
    seen.add(id(el))
    body.append(el)
if sect_pr is not None:
    body.append(sect_pr)

doc.save(OUT)
print(f"saved: {OUT}")
print(f"elements: {len(seen)}")
