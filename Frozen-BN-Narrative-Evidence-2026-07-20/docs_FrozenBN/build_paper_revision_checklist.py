#!/usr/bin/env python3
"""Build the paper revision checklist as a Word document.

Every item: WHERE (estimated page + exact words to Cmd+F), PROBLEM (what is
wrong and why), FIX (what it needs to become -- guidance, user writes own
words), REF (citation + link + where inside the source).

Run:
  /Library/Frameworks/Python.framework/Versions/3.11/bin/python3.11 \
      docs/build_paper_revision_checklist.py
Writes docs/paper_revision_checklist.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches

OUT = "docs/paper_revision_checklist.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

doc.add_heading("NTSB Paper Revision Checklist", level=0)
p = doc.add_paragraph()
p.add_run("Draft reviewed: Master Draft NTSB Paper 6_19_26.docx  |  "
          "Checklist date: July 21, 2026").italic = True
doc.add_paragraph(
    "How to use: work top to bottom. For each item, Cmd+F the exact quoted "
    "words in your draft (page numbers are estimates; the quote is the "
    "reliable locator). Rewrite in YOUR OWN words, then send the section "
    "back for review. Check the box when done.")

GLOBAL = [
    ("G1. Citation style",
     "The draft cites author-year (\"Zhang and Mahadevan, 2021\"). RESS / "
     "Elsevier style is numbered references [1], [2] in order of first "
     "appearance, and every website citation ends with \"Accessed [date]\". "
     "Convert everything, including the reference list."),
    ("G2. \"Incidents\" vs \"accidents\"",
     "NTSB formally distinguishes accidents from incidents. The 2,243 "
     "FAR-121 records are ACCIDENTS (Zhang's word too). Global "
     "find-and-replace, then re-read each spot."),
    ("G3. Dates",
     "Every \"1982-2016\" becomes 1982-2019 (verified against the dataset: "
     "2,243 accidents, 1982-2019). Appears in Abstract, Introduction, and "
     "Section 3. Also \"1982-2026\" in the Introduction is a typo."),
    ("G4. Lost math symbols",
     "Every equation lost its Greek/math symbols: summation shows as \"_j\" "
     "or \"_K\"; alpha is blank in \"with  = 2.0\" and \"exp(  x structure "
     "similarity)\" (Sec 7.6). Rebuild all equations in Word's equation "
     "editor and number them (1), (2), (3)... like Zhang."),
    ("G5. Missing core sections",
     "The paper ends at the old retrieval pipeline. Missing entirely: the "
     "Zhang BN reproduction, the narrative-to-evidence bridge "
     "(deterministic parser + LLM tier, hard/soft/stated evidence), and the "
     "296-accident held-out validation with significance stats. Per the "
     "Zhang-style outline these become methodology 4.1/4.3/4.4 and results "
     "5.1-5.5."),
]

# (section header, [(where, problem, fix, ref), ...])
ITEMS = [
 ("Title page + Abstract (p. ~1)", [
  ("\"aviation incident reports that happened from 1982-2016\"",
   "Wrong dates (G3); \"that happened\" is filler.",
   "1982-2019; tighten the sentence.", ""),
  ("\"Two similar methodologies are created\"",
   "Diagnosis and prognosis are two MODES of one methodology, not two "
   "methodologies. Bigger problem: the abstract only describes the "
   "counting path -- never mentions the Bayesian network, the parser, or "
   "the held-out validation (the actual contribution).",
   "Rewrite the abstract LAST, after the body is done, covering: "
   "reproduction of Zhang, narrative bridge, held-out results.", ""),
  ("\"predict the downstream causes of incidents and future forthcoming "
   "events\"",
   "Causes are upstream (diagnosis); events are downstream (prognosis). "
   "\"Downstream causes\" mixes them. \"Future forthcoming\" is redundant.",
   "Say upstream causes and downstream (next) events.", ""),
  ("\"a streamlit application, an interactive tool in the Python library\"",
   "Garbled phrase.",
   "Streamlit is a Python library for building interactive web apps.", ""),
 ]),
 ("Section 1: Introduction (p. ~1-2)", [
  ("\"The Federal Aviation Agency (FAA)\"",
   "It is the Federal Aviation ADMINISTRATION. First-paragraph error a "
   "reviewer catches instantly.",
   "Fix the name.", ""),
  ("\"ensure global economic stability. (Ref)\"",
   "Placeholder reference.",
   "Cite FAA Air Traffic by the Numbers; headline stats (daily flights, "
   "passengers) are on the landing page.",
   "https://www.faa.gov/air_traffic/by_the_numbers"),
  ("\"businesses across the world. (re)\"",
   "Placeholder reference. RESOLVED during review: this sentence is "
   "general framing and needs NO citation -- just delete the (re). Do "
   "NOT use Zhang's IATA press-release link (his ref [1]): it is dead "
   "(404, verified July 21, 2026). The Safety Report landing page alone "
   "does not show figures either (they are behind its Executive Summary "
   "link).",
   "Delete the (re). The paragraph's IATA statistic (5.1 billion "
   "passengers worldwide in 2026) cites the working June 2026 press "
   "release instead; the figure is in the bulleted Highlights list near "
   "the top: \"Passenger numbers are expected to reach 5.1 billion in "
   "2026 (up 2.4% on 2025).\"",
   "https://www.iata.org/en/pressroom/2026-releases/06-07-middle-east-"
   "disruptions-high-fuel-prices-halve-airline-industry-profitability/"),
  ("\"The NTSB consisted of commercial aviation data from 1982-2026\"",
   "\"2026\" is a typo; the AGENCY did not consist of data -- the DATABASE "
   "contains it. Also \"Their mission was\" should be \"is\" (agency still "
   "exists).",
   "The NTSB database contains...; mission is...; add the database "
   "citation.",
   "https://data.ntsb.gov/avdata (the downloadable dataset page)"),
  ("\"2,243 incidents from FAR 121, between 1982-2016\"",
   "Accidents (G2), 1982-2019 (G3), and \"FAR 121\" is used without a "
   "citation.",
   "Fix terms and cite the FAA Part 121 page (defines Part 121 carriers). "
   "This is Zhang's ref [31].",
   "https://www.faa.gov/hazmat/air_carriers/operations/part_121/"),
  ("\"Earlier Bayesian network analyses relied on the coded fields rather "
   "than the text narratives.\"",
   "This sentence is about Zhang; it needs his citation right here.",
   "Cite Zhang & Mahadevan 2021; his Section 2 describes using the coded "
   "event-sequence data.",
   "https://doi.org/10.1016/j.ress.2020.107371"),
  ("\"creates new opportunities to utilize the unstructured text with "
   "efficiency. (re)\"",
   "Placeholder reference.",
   "Cite an NLP-for-aviation-safety paper (read the abstracts first).",
   "Tanguy et al. 2016: https://doi.org/10.1016/j.compind.2015.09.005  "
   "and/or Rose et al. 2020: https://doi.org/10.3390/aerospace7100143"),
  ("\"Event trees are contradictory methods to Fault trees.\"",
   "Wrong word: they are COMPLEMENTARY (fault trees reason backward to "
   "causes; event trees forward to outcomes). Your own next sentences say "
   "this, so \"contradictory\" contradicts your own text.",
   "Say complementary; cite a reliability textbook (fault trees and event "
   "trees each have a chapter).",
   "Rausand & Hoyland, System Reliability Theory: "
   "https://doi.org/10.1002/9780470316900 (verify link loads)"),
  ("Two Bayesian-network paragraphs (\"They are probabilistic graphical "
   "models...\" and \"Bayesian networks are probabilistic models that "
   "represent relationships...\")",
   "The same definition appears twice in back-to-back paragraphs.",
   "Merge into one paragraph; cite Pearl (Ch. 3 introduces belief "
   "networks) and Jensen 1996 (Zhang's ref [39]).",
   "Pearl 1988: https://doi.org/10.1016/C2009-0-27609-4 (verify link)"),
  ("\"LLMs probabilities are not grounded in data\" (and the 3 bullets)",
   "Grammar (LLMs' or \"LLM probabilities\") and conversational tone "
   "(\"It's too confident...\").",
   "Keep the content, formalize the tone; cite the hallucination claim.",
   "Ji et al. 2023, ACM Computing Surveys: https://doi.org/10.1145/3571730 "
   "(abstract + Sec 1 define hallucination)"),
  ("MISSING at end of Introduction",
   "Zhang's intro ends with a numbered contribution list and a roadmap "
   "paragraph (\"Section 2 describes...\"). Yours has neither.",
   "Add both. Contributions: reproduction of Zhang, the narrative bridge, "
   "held-out validation.", ""),
 ]),
 ("Section 2: Related Work (p. ~2-3)", [
  ("2.1 first sentence \"The inspiration for this alternative... pioneered "
   "by Zhang and Mahadevan, 2021.\"",
   "Word-for-word duplicate of the same sentence in the abstract.",
   "Rewrite one of the two.", ""),
  ("\"prior possible probabilities were estimated\"",
   "\"Possible\" doesn't belong.",
   "\"prior probabilities were estimated\"", ""),
  ("\"for aligning the casual chains\"",
   "Typo: CAUSAL, not casual. Search the whole paper -- this typo tends to "
   "repeat.",
   "Fix everywhere.", ""),
  ("\"we retrieved and reranked our embeddings from Nogueira and Cho "
   "(2019)\"",
   "You didn't get embeddings from them; you adopted their RERANKING idea.",
   "Say the reranking approach follows Nogueira & Cho.",
   "Gentner: https://doi.org/10.1207/s15516709cog0702_3 | Falkenhainer: "
   "https://doi.org/10.1016/0004-3702(89)90077-5 | Goldstone: "
   "https://doi.org/10.1037/0278-7393.20.1.3 | Needleman & Wunsch: "
   "https://doi.org/10.1016/0022-2836(70)90057-4 | Nogueira & Cho: "
   "https://arxiv.org/abs/1901.04085"),
  ("\"Spencer-Smith and Goldstone, 1997\"",
   "Could not be verified as a real citation.",
   "ASK JESSE SPENCER-SMITH directly for the exact citation -- he is a "
   "co-author. Do not leave it in unverified.", ""),
  ("MISSING paragraphs in Related Work",
   "A RESS reviewer expects breadth: BNs in transportation safety and ML "
   "for aviation risk.",
   "Add one paragraph on BN safety applications (reuse from Zhang's list: "
   "Ale et al. CATS [3][22][34], Luxhoj 2003 [21], Ancel et al. 2015 "
   "[29]) and one on ML baselines (Zhang & Mahadevan 2019, Decision "
   "Support Systems -- pairs with your LR baseline). Read abstracts before "
   "citing.", ""),
 ]),
 ("Section 3: NTSB Accident Investigation Data (p. ~3)", [
  ("\"2,243 incidents from the time period between 1982-2016\"",
   "G2 + G3.",
   "2,243 accidents, 1982-2019.", ""),
  ("\"collected, scraped, and retrieved from Zhang's GitHub for a total of "
   "14 files\"",
   "Repo needs a citation with access date. Also inconsistent file count: "
   "here 14 files (8 used), Sec 5.1 says \"look at all 14 files\", Sec 6.1 "
   "says \"eight different files\".",
   "Cite the repo URL (verify it is still public); pick one consistent "
   "file-count story.", ""),
  ("\"sequence of events for each incident that had recovered\"",
   "Garbled -- \"that had recovered\" means nothing here.",
   "\"that were recorded\".", ""),
  ("MISSING: train/test split definition",
   "The split is load-bearing for the whole validation story and appears "
   "nowhere.",
   "State: 1,742 accidents in the 1982-2006 window (matches Zhang); 296 "
   "held-out accidents 2007-2019 with usable narratives.", ""),
  ("MISSING: coding manual citation where the code dictionary is described",
   "The coded vocabulary has an official source.",
   "Cite the NTSB Aviation Coding Manual (Zhang's ref [35]) -- it defines "
   "the coded vocabulary your parser maps into.",
   "https://www.ntsb.gov/GILS/Documents/codman.pdf"),
 ]),
 ("Section 4: Background Concepts (p. ~4)", [
  ("4.1: \"the value near1 means\" / \"multi-dimensional mode\" / \"range "
   "goes between 0 (unrelated) to 1(identical)\"",
   "Missing space; \"mode\" should be \"space\"; cosine similarity ranges "
   "-1 to 1 (in practice ~0 to 1 with these embeddings).",
   "Fix all three; cite the OpenAI embeddings docs (model name and "
   "dimensions in the models table).",
   "https://platform.openai.com/docs/guides/embeddings"),
  ("4.1 equation: \"sim (q, i) = (q, i) / (||q|| x ||i||)\"",
   "Numerator must be the dot product.",
   "q . i in the equation editor; number the equation (G4).", ""),
  ("4.2: \"((https://www.ibm.com/think/topics/k-means-clustering...))\"",
   "A marketing page is not citable in a journal.",
   "Delete the URL; cite MacQueen 1967 (original k-means paper, on "
   "Project Euclid -- search \"MacQueen Some methods for classification\") "
   "and scikit-learn for the implementation.",
   "https://jmlr.org/papers/v12/pedregosa11a.html"),
  ("4.3: \"Each weight is weighted by how likely that scenario is\"",
   "Circular (\"weight is weighted\").",
   "Each CONDITIONAL PROBABILITY is weighted by the scenario's "
   "probability. Fix the missing summation symbol (G4).", ""),
  ("4.4: \"a dynamic method\"",
   "Imprecise.",
   "\"dynamic-programming algorithm\".", ""),
  ("MISSING: two background subsections",
   "The new core sections need them.",
   "(a) Bayesian networks (move merged intro material here; cite "
   "Pearl/Jensen). (b) Soft evidence / Jeffrey conditioning -- cite "
   "Jeffrey, The Logic of Decision, Chapter 11 \"Probability kinematics\" "
   "(that chapter is where the method lives).", ""),
 ]),
 ("Section 5: Initial LLM-Only Approach (p. ~4-5)", [
  ("Orphan paragraph: \"The next step is retrieval. We utilize an LLM and "
   "implement prompting to look at all 14 files...\"",
   "Repeats the retrieval step described two paragraphs earlier, "
   "contradicts the 8-file count, and sits after the prognosis paragraph "
   "out of order.",
   "Delete or merge it.", ""),
  ("5.2: \"Changing a single word can change the whole pipeline\"",
   "The pipeline doesn't change; the OUTPUT does.",
   "Say the output changes.", ""),
  ("5.3: \"the LLM implemented the scoring and produced the probabilities "
   "itself but were ungrounded\"",
   "Grammar -- singular subject, plural verb.",
   "Fix agreement.", ""),
  ("5.3: \"so we were not getting hallucinated data that we could rely "
   "on\"",
   "As written this says hallucinated data is reliable -- the opposite of "
   "your point.",
   "You mean: probabilities had to be grounded in data so results are "
   "reproducible and trustworthy.", ""),
 ]),
 ("Section 6: Proposed Methodology (p. ~5-6)", [
  ("6.1: \"a data column that represents a particular incident\"",
   "Imprecise.",
   "EV_ID is the unique identifier (key) for an accident.", ""),
  ("6.2: \"text_embedding_3_small\" (both places)",
   "Wrong model name format.",
   "text-embedding-3-small (hyphens, matching OpenAI).", ""),
  ("6.4: \"the results show clusters failure themes for that query\" / "
   "\"named from the LLM\"",
   "Garbled; \"from\" should be \"by\".",
   "Fix both; VERIFY \"50 clusters\" against the code before submitting.",
   ""),
  ("6.5: \"then the diagnosis part happens\"",
   "Conversational.",
   "Describe the computation; number the formulas (G4).", ""),
 ]),
 ("Section 7: Structural Mapping (p. ~6-8)", [
  ("7.1: \"but actually are matching by the causes\"",
   "Grammar.",
   "Fix.", ""),
  ("7.3: \"The query causal cause is also extracted\"",
   "Wrong phrase.",
   "\"query causal chain\".", ""),
  ("7.5: \"The lengths of these causal chains have different lengths\"",
   "Redundant.",
   "\"These causal chains have different lengths.\"", ""),
  ("7.6: \"so the new score equals cosine A2(structural mapping) reduces "
   "exactly to A0\"",
   "Two sentences fused; alpha symbols missing (G4).",
   "Separate the sentences; restore alpha everywhere.", ""),
  ("STRUCTURAL DECISION for Sections 7-9",
   "Under the Zhang-style outline these compress into one results "
   "subsection (ablation) or an appendix. The finding was null (McNemar "
   "p = 1.000), so it earns a paragraph plus a table, not three sections.",
   "Keep the honest content (\"tested, neither helps nor hurts\"), "
   "shorter.", ""),
 ]),
 ("Section 8: Worked Examples (p. ~8-11)", [
  ("ALL numbers in the tables (53.5%, 18.1%, 40.7%, ...)",
   "They came from the old pipeline; the parser and pipeline have changed "
   "since.",
   "Re-run the examples before submitting (ask for regeneration when you "
   "reach this section) or the tables won't match the released code.", ""),
  ("8.2.1: \"to show the method's behavior is not strong and it depends on "
   "the query\"",
   "Garbled and needlessly self-damaging.",
   "What the case shows: retrieval is sensitive to query style; long "
   "multi-topic narratives dilute retrieval; the focused query recovers "
   "it.", ""),
  ("MISSING: the main worked example",
   "Under the new structure the primary example is the real held-out 2007 "
   "accident walkthrough (narrative -> parser tier -> evidence with "
   "strengths -> network probabilities).",
   "Source: outputs/real_narrative_walkthrough.md. Keep the old "
   "engine-fire / landing-gear examples as counting-path illustrations.",
   ""),
 ]),
 ("Section 9: Aggregate Validation (p. ~11-12)", [
  ("\"Figure 9.2. Paired flip table\"",
   "It's a table, not a figure.",
   "Call it Table 9.2 (renumber in final structure).", ""),
  ("\"but further testing with larger sample size can be done to test this "
   "theory further\"",
   "Redundant double \"test/further\".",
   "Tighten.", ""),
  ("\"254 incidents... 177 to train and 77 are held out\"",
   "Unverified numbers.",
   "Verify against the code before submitting.", ""),
  ("MISSING: the primary validation (new)",
   "The 296-accident held-out evaluation is now the headline validation. "
   "LEAK-SAFE numbers only (outcome phrases redacted before embedding; "
   "stated-severity readout OFF) -- the older 92.9%/81.4% figures were "
   "pre-audit and read severity wording from the text; do NOT use them.",
   "Numbers for the table: bn-sev (k-NN severity as virtual evidence "
   "through the frozen BN) 90.9% injury [87.5, 93.9] / 77.4% damage "
   "[72.3, 82.1]; BN event path (soft-priority) 89.9% / 55.7%; LR on "
   "parsed features 87.8% / 64.2%; LR on narrative embedding 91.6% / "
   "74.0%. bn-sev beats parsed-feature LR at McNemar p = 0.0225 (injury), "
   "p < 0.0001 (damage); vs embedding LR not significant (p = 0.50 / "
   "0.11); vs no-narrative prior p < 0.0001 on both. Severe-outcome "
   "screening: injury 93.5% sensitivity / 96.8% specificity. Source: "
   "outputs/heldout_significance.md. Cite Brier 1950 where the Brier "
   "score is defined, and McNemar 1947 (Psychometrika 12(2):153-157).",
   "Brier 1950: https://doi.org/10.1175/1520-0493(1950)078<0001:VOFEIT>"
   "2.0.CO;2"),
 ]),
 ("Section 10: Discussion (p. ~12-13)", [
  ("10.1: \"the specific causes that are given are not fully current and "
   "the probabilities are low - which would not make pilots act on these "
   "probabilities\"",
   "\"Not fully current\" is garbled (you mean not precise enough / spread "
   "thin), and the sentence concludes your own method is unusable. It "
   "reflects only the old 32.5% retrieval result.",
   "Rewrite the conclusion around the full story: reproduction verified, "
   "narratives add significant signal, held-out accuracy beats a "
   "supervised baseline.", ""),
  ("10.2: \"Using unstructured data as context to find more information to "
   "find more accurate results to find out what causes these incidents.\"",
   "Three \"to find\"s and it's a sentence fragment.",
   "Mirror Zhang's 6.1: a NUMBERED list of contributions.", ""),
  ("10.3: \"LLM extraction errors propagate into structural similarity "
   "which is a misleading step in the structural score\"",
   "Garbled.",
   "You mean extraction errors can corrupt the structural score. Also add "
   "new limitations: single dataset (FAR 121); LLM tier's measured "
   "accuracy below the deterministic path; damage prediction via the BN "
   "chain not significantly better than LR.", ""),
  ("10.4: \"having an LLM take the narratives and turn them into coded "
   "data and then have an LLM create a skill to create a Bayesian "
   "network... and compare from Zhang's and Mahadevan work\"",
   "THIS IS COMPLETED WORK, NOT FUTURE WORK. It is the heart of the paper "
   "now.",
   "Remove from Future Work; replace with real future items: other "
   "datasets (ASRS, other FAR parts), richer severity targets, "
   "prospective evaluation.", ""),
  ("MISSING after Section 10: Declarations + References",
   "Required by the journal.",
   "Add: AI-use declaration (agreed wording), data availability (Zhang's "
   "GitHub + your repo), acknowledgments if any, then the numbered "
   "reference list in order of first appearance.", ""),
 ]),
]


def add_item(n, where, problem, fix, ref):
    p = doc.add_paragraph()
    r = p.add_run(f"[  ]  Item {n}")
    r.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.add_run("Where: ").bold = True
    p.add_run(where)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.add_run("Problem: ").bold = True
    p.add_run(problem)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.add_run("Fix (your words): ").bold = True
    p.add_run(fix)
    if ref:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.add_run("Reference / link: ").bold = True
        r = p.add_run(ref)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x9D)


doc.add_heading("Global fixes (apply everywhere first)", level=1)
for title, body in GLOBAL:
    p = doc.add_paragraph()
    p.add_run(f"[  ]  {title}. ").bold = True
    p.add_run(body)

n = 0
for section, items in ITEMS:
    doc.add_heading(section, level=1)
    for where, problem, fix, ref in items:
        n += 1
        add_item(n, where, problem, fix, ref)

doc.add_heading("Part 2: New sections to WRITE (not fixes -- new content)",
                level=1)
doc.add_paragraph(
    "The fixes above repair what exists. These are the sections that do "
    "not exist yet and must be written from scratch. Target structure "
    "mirrors Zhang's paper (1 Introduction, 2 Data, 3 Background, "
    "4 Proposed methodology, 5 Computational results, 6 Discussion).")

NEW_SECTIONS = [
    ("4.1 Reproducing Zhang's Bayesian network construction",
     "His Section 4 recipe (priors, edge selection, Beta-CDF CPTs), your "
     "implementation in pyAgrum, and the two upgrades (person-finding "
     "nodes, multi-state severity nodes).",
     "Source: tests/bn_upgraded.py, docs/ALL_TABLES_EXACT_COMPARISON.md"),
    ("4.3 The narrative-to-evidence parser (two tiers)",
     "Deterministic parser (exact NTSB vocabulary, confidence 1.0) runs "
     "first; LLM tier (gpt-4.1) only when no event facts found; every "
     "LLM-proposed fact gets its strength measured from the data, never "
     "from the LLM's own confidence.",
     "Source: llm_evidence.py, outputs/combined_parser_validation.md"),
    ("4.4 Evidence types and how they enter the network",
     "Hard evidence (named vocabulary, set directly), soft evidence "
     "(described facts, entered via virtual evidence / Jeffrey "
     "conditioning), stated severity (outcomes mentioned in text, "
     "weighted by historical reliability).",
     "Source: the architecture figure "
     "docs/figures/narrative_to_bn_architecture.png"),
    ("5.1 Reproduction results (scoreboard)",
     "Counting layer: Table 7, 85/85 exact. Network layer: 93 published "
     "values -> 48 exact / 29 close / 12 differ / 4 qualitative after "
     "upgrades; the 12 differences traced via the 30-seed variance "
     "envelope, and on the 4 queryable cells your build is as close or "
     "closer to the published values than Zhang's own released model "
     "file.",
     "Source: docs/ALL_TABLES_EXACT_COMPARISON.md, "
     "outputs/BN_UPGRADED_ENVELOPE.md, docs/TABLE7_FULL_REPRODUCTION.md"),
    ("5.2 Parser validation",
     "Identity (11/11 scoreboard sentences, tier 1 only), safety (6 "
     "adversarial probes), paraphrase recovery (tier 2 vs retrieval-only "
     "fallback).",
     "Source: outputs/combined_parser_validation.md"),
    ("5.3 Held-out evaluation (the headline table)",
     "296 unseen accidents 2007-2019; all predictors with 95% CIs; "
     "McNemar pairwise tests; the LR baseline comparison; honest caveats "
     "(damage via BN chain ties LR; LLM tier below deterministic path on "
     "injury).",
     "Source: outputs/heldout_significance.md, "
     "outputs/heldout_narrative_bn_eval.json"),
    ("5.4 Walkthrough on a real held-out accident",
     "One 2007 accident end to end: narrative -> tier decision -> "
     "evidence with strengths -> network probabilities vs what actually "
     "happened. Mirrors Zhang's scenario-analysis section.",
     "Source: outputs/real_narrative_walkthrough.md"),
    ("5.5 Ablations",
     "LLM-first (LLM as sole front door) vs tiered; retrieval-only "
     "fallback vs LLM tier; the old structural-mapping A0/A2 result "
     "compressed to a paragraph + table here.",
     "Source: outputs/heldout_significance.md (llm-first rows), "
     "outputs/llm_only_identity.md, draft Sections 7-9"),
    ("Declarations block (after Discussion)",
     "AI-use declaration, data availability, acknowledgments, then the "
     "numbered reference list.",
     "Source: agreed AI-use wording from review discussion"),
]
for title, what, src in NEW_SECTIONS:
    p = doc.add_paragraph()
    p.add_run(f"[  ]  {title}. ").bold = True
    p.add_run(what + "  ")
    r = p.add_run(src)
    r.italic = True
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_heading("Part 3: Reference map -- [n], where it goes, link, and "
                "what the source says", level=1)
doc.add_paragraph(
    "Numbers are assigned in order of FIRST appearance (Zhang/Elsevier "
    "style). If your final sentence order differs, renumber -- the "
    "assignments below assume the checklist structure. Every web citation "
    "gets \"Accessed July 2026\" (or your actual access date). Quotes "
    "marked VERIFIED were fetched and checked on July 21, 2026; the rest "
    "describe what you will find -- open each link and confirm before "
    "citing.")

REFS = [
 ("[1] FAA, Air Traffic by the Numbers",
  "Introduction, first paragraph -- replaces \"(Ref)\" after \"ensure "
  "global economic stability\".",
  "https://www.faa.gov/air_traffic/by_the_numbers",
  "VERIFIED -- the page opens with this exact sentence: \"Every day, "
  "FAA's Air Traffic Organization (ATO) provides service to more than "
  "44,000 flights and more than 3 million airline passengers across more "
  "than 29 million square miles of airspace.\" Use those figures for the "
  "scale claim. (FAA blocks some automated access; opens fine in a "
  "browser.)"),
 ("[2] IATA press release, June 2026 (passenger statistic)",
  "Introduction, first paragraph -- the sentence with the 5.1 billion "
  "passengers worldwide statistic. NOTE: as written this is the FIRST "
  "citation in the paper, so it takes number [1]; renumber accordingly.",
  "https://www.iata.org/en/pressroom/2026-releases/06-07-middle-east-"
  "disruptions-high-fuel-prices-halve-airline-industry-profitability/",
  "VERIFIED July 21, 2026 -- the bulleted Highlights list near the top "
  "states: \"Passenger numbers are expected to reach 5.1 billion in "
  "2026 (up 2.4% on 2025).\" WARNING: Zhang's own IATA citation (his "
  "ref [1], a 2018 press release) is a DEAD LINK (404) -- do not copy "
  "it. The IATA Annual Safety Report landing page "
  "(https://www.iata.org/en/publications/safety-report/) is also NOT a "
  "good citation for figures -- data sits behind its Executive Summary "
  "link."),
 ("[3] NTSB Aviation Accident Database",
  "Introduction, second paragraph (the NTSB description) and again in "
  "Section 3 (data).",
  "https://data.ntsb.gov/avdata",
  "VERIFIED -- the page is the NTSB download directory; you will see "
  "avall.zip (the complete aviation accident database, updated monthly) "
  "and codman.pdf listed. This is the official source of the data Zhang "
  "scraped."),
 ("[4] FAA, Regularly Scheduled Air Carriers (Part 121)",
  "Introduction / Section 3, the sentence \"2,243 accidents from FAR "
  "121\" -- cite at first mention of Part 121.",
  "https://www.faa.gov/hazmat/air_carriers/operations/part_121/",
  "Could not be fetched automatically (FAA 403) -- open in browser. The "
  "page defines Part 121 operators: scheduled U.S. airlines, regional "
  "air carriers, and cargo operators. Zhang's paper words it: \"U.S.-based "
  "airlines, regional air carriers, and all cargo operators [31]\". This "
  "is Zhang's ref [31]."),
 ("[5] Zhang X, Mahadevan S. Bayesian network modeling of accident "
  "investigation reports for aviation safety assessment. Reliab Eng Syst "
  "Saf 2021;209:107371.",
  "Introduction (\"Earlier Bayesian network analyses relied on the coded "
  "fields...\"), Related Work 2.1, and throughout methodology whenever "
  "you reference his recipe.",
  "https://doi.org/10.1016/j.ress.2020.107371",
  "His Section 2 describes the NTSB coded data; Section 4 is the "
  "construction recipe (4.2 priors, 4.3 conditional probabilities, 4.4 "
  "network construction); Section 5 has the tables you reproduce "
  "(Tables 7-9, Figs 11-12). You have the PDF: docs/BN-NTSB RESS "
  "2021.pdf."),
 ("[6] Tanguy L, Tulechki N, Urieli A, Hermann E, Raynal C. Natural "
  "language processing for aviation safety reports. Comput Ind 2016.",
  "Introduction, the sentence about LLMs/NLP creating \"new opportunities "
  "to utilize the unstructured text\" -- and Related Work.",
  "https://doi.org/10.1016/j.compind.2015.09.005",
  "Abstract describes applying NLP to large corpora of aviation safety "
  "reports (classification and analysis). READ THE ABSTRACT before "
  "citing to confirm it supports your sentence."),
 ("[7] Rose R, Puranik T, Mavris D. Natural language processing based "
  "method for clustering and analysis of aviation safety narratives. "
  "Aerospace 2020;7(10):143.",
  "Same spot as [6] (second citation), or Related Work.",
  "https://doi.org/10.3390/aerospace7100143",
  "Open access (MDPI) -- abstract describes clustering aviation safety "
  "narratives with NLP. Verify abstract."),
 ("[8] Rausand M, Hoyland A. System Reliability Theory: Models, "
  "Statistical Methods, and Applications. Wiley.",
  "Introduction, the fault tree / event tree paragraph.",
  "https://doi.org/10.1002/9780470316900",
  "Standard reliability textbook; fault tree analysis and event tree "
  "analysis each have a dedicated chapter (check the table of contents "
  "on the Wiley page). Verify the link resolves; if not, cite the "
  "edition you can access."),
 ("[9] Pearl J. Probabilistic Reasoning in Intelligent Systems: Networks "
  "of Plausible Inference. Morgan Kaufmann/Elsevier.",
  "Introduction / Background, the merged Bayesian-network paragraph; "
  "also where virtual evidence is described in the new methodology.",
  "https://doi.org/10.1016/C2009-0-27609-4",
  "THE standard BN citation (it is Zhang's ref [36] too -- he cites the "
  "Elsevier 2014 reprint). Chapter 3 introduces belief networks. Verify "
  "the DOI resolves; otherwise copy Zhang's exact reference entry."),
 ("[10] Jensen FV. An Introduction to Bayesian Networks. UCL Press; "
  "1996.",
  "Same paragraph as [9], second citation.",
  "(book -- no stable link; this is Zhang's ref [39], copy his entry)",
  "Introductory BN textbook. Cite alongside Pearl."),
 ("[11] Ji Z, et al. Survey of hallucination in natural language "
  "generation. ACM Comput Surv 2023.",
  "Introduction, the LLM-limitations bullets (\"LLMs probabilities are "
  "not grounded in data, and they hallucinate\").",
  "https://doi.org/10.1145/3571730",
  "Abstract + Section 1 define hallucination as generated content that "
  "is nonsensical or unfaithful to the source. Supports your grounding "
  "argument. Open access on arXiv if the DOI is paywalled."),
 ("[12] Zhang X, Mahadevan S. Ensemble machine learning models for "
  "aviation incident risk prediction. Decis Support Syst 2019;116:48-63.",
  "Related Work, the new ML-for-aviation paragraph -- pairs with your "
  "supervised LR baseline.",
  "https://doi.org/10.1016/j.dss.2018.10.009",
  "This is Zhang's ref [2]. Abstract: supervised ML models predicting "
  "aviation incident risk. Verify abstract; if the DOI does not resolve, "
  "copy his reference entry."),
 ("[13] Ale BJ, et al. Further development of a causal model for air "
  "transport safety (CATS): building the mathematical heart. Reliab Eng "
  "Syst Saf 2009;94(9):1433-41.",
  "Related Work, the new BN-in-transportation-safety paragraph.",
  "https://doi.org/10.1016/j.ress.2009.02.024",
  "Zhang's ref [3] -- BN-based causal model for air transport safety. "
  "Verify the DOI (copy his entry if needed)."),
 ("[14] Luxhoj JT. Probabilistic causal analysis for system safety risk "
  "assessments in commercial air transport. IRIA 2003.",
  "Same paragraph as [13].",
  "(workshop paper -- search Google Scholar for the PDF)",
  "Zhang's ref [21]. BN risk assessment in commercial air transport."),
 ("[15] Ancel E, et al. Predictive safety analytics: inferring aviation "
  "accident shaping factors and causation. J Risk Res "
  "2015;18(4):428-51.",
  "Same paragraph as [13].",
  "https://doi.org/10.1080/13669877.2014.896402",
  "Zhang's ref [29] -- object-oriented BN inferring accident shaping "
  "factors. Verify DOI."),
 ("[16] Gentner D. Structure-mapping: a theoretical framework for "
  "analogy. Cogn Sci 1983;7(2):155-70.",
  "Section 2.2 (Structural Mapping Logic), first sentence.",
  "https://doi.org/10.1207/s15516709cog0702_3",
  "Abstract: analogy is a mapping of RELATIONS between objects, not "
  "attributes -- the basis of your role/system/mechanism comparison."),
 ("[17] Falkenhainer B, Forbus KD, Gentner D. The structure-mapping "
  "engine: algorithm and examples. Artif Intell 1989;41(1):1-63.",
  "Section 2.2, second sentence.",
  "https://doi.org/10.1016/0004-3702(89)90077-5",
  "Sections 1-2 describe the SME algorithm that scores alignments on "
  "structural evidence -- what your chain-alignment scoring is modeled "
  "on."),
 ("[18] Goldstone RL. Similarity, interactive activation, and mapping. "
  "J Exp Psychol Learn Mem Cogn 1994;20(1):3-28.",
  "Section 2.2, the sentence about similarity and structural mapping "
  "interacting.",
  "https://doi.org/10.1037/0278-7393.20.1.3",
  "Abstract: similarity assessments and mapping interact rather than "
  "being independent."),
 ("[19] Spencer-Smith & Goldstone, 1997",
  "Section 2.2, the weighting-matched-features sentence.",
  "COULD NOT BE VERIFIED",
  "ASK JESSE SPENCER-SMITH for the exact citation -- he is a co-author "
  "on your paper. Do not submit unverified."),
 ("[20] Needleman SB, Wunsch CD. A general method applicable to the "
  "search for similarities in the amino acid sequence of two proteins. "
  "J Mol Biol 1970;48(3):443-53.",
  "Section 2.2 and Section 4.4 / 7.5 (chain alignment).",
  "https://doi.org/10.1016/0022-2836(70)90057-4",
  "The short paper presents the dynamic-programming alignment algorithm "
  "with gaps -- exactly what you apply to causal chains."),
 ("[21] Nogueira R, Cho K. Passage re-ranking with BERT. arXiv "
  "2019;1901.04085.",
  "Section 2.2 (the reranking sentence) and Section 7.6 (score fusion).",
  "https://arxiv.org/abs/1901.04085",
  "Free on arXiv. Abstract: rerank retrieved passages with a neural "
  "model -- the retrieve-then-rerank pattern your A2 follows."),
 ("[22] Zhang's data repository (GitHub)",
  "Section 3, \"collected, scraped, and retrieved from Zhang's GitHub\".",
  "(you have the URL -- paste it and confirm it is still public)",
  "Cite as a web reference with access date."),
 ("[23] NTSB Aviation Coding Manual",
  "Section 3, where the code dictionary is described; also the new "
  "parser section (the vocabulary source).",
  "https://www.ntsb.gov/GILS/Documents/codman.pdf",
  "VERIFIED to exist (also listed as codman.pdf on the data.ntsb.gov/"
  "avdata page). It is the official manual of NTSB occurrence/subject "
  "codes -- the vocabulary your deterministic parser matches against. "
  "This is Zhang's ref [35]."),
 ("[24] OpenAI, text embeddings documentation",
  "Section 4.1 (embeddings background) and 6.2 (model name).",
  "https://platform.openai.com/docs/guides/embeddings",
  "VERIFIED -- the page states: \"By default, the length of the "
  "embedding vector is 1536 for text-embedding-3-small\" (your "
  "dimension claim), \"We recommend cosine similarity\", and \"OpenAI "
  "embeddings are normalized to length 1\" (why cosine equals dot "
  "product here). Model name is text-embedding-3-small with hyphens."),
 ("[25] MacQueen J. Some methods for classification and analysis of "
  "multivariate observations. Proc 5th Berkeley Symp 1967:281-97.",
  "Section 4.2 -- REPLACES the IBM URL currently in the draft.",
  "https://projecteuclid.org/ebooks/berkeley-symposium-on-mathematical-"
  "statistics-and-probability (search \"MacQueen\")",
  "The original k-means paper; Section 1 introduces the k-means "
  "procedure of partitioning by nearest centroid."),
 ("[26] Pedregosa F, et al. Scikit-learn: machine learning in Python. "
  "J Mach Learn Res 2011;12:2825-30.",
  "Section 4.2 (k-means implementation) and the new held-out section "
  "(LR baseline implementation).",
  "https://jmlr.org/papers/v12/pedregosa11a.html",
  "Free PDF at that link -- the standard citation for the scikit-learn "
  "library you used for KMeans and LogisticRegression."),
 ("[27] Jeffrey RC. The Logic of Decision. 2nd ed. University of "
  "Chicago Press; 1983.",
  "New background subsection on soft evidence, and the new methodology "
  "section where described facts enter as virtual evidence.",
  "(book -- search press.uchicago.edu for \"The Logic of Decision\")",
  "Chapter 11, \"Probability kinematics\", is where updating on "
  "uncertain evidence (Jeffrey conditioning) is defined -- the formal "
  "basis for your soft-evidence entry."),
 ("[28] Ducamp G, Gonzales C, Wuillemin P-H. aGrUM/pyAgrum: a toolbox "
  "to build models and algorithms for probabilistic graphical models in "
  "Python. PGM 2020.",
  "New methodology section 4.1, first mention of the BN implementation.",
  "https://proceedings.mlr.press/v138/ducamp20a.html",
  "Free PDF on the PMLR page -- the citation for the pyAgrum library "
  "your network is built in. Verify the page loads."),
 ("[29] BayesFusion, GeNIe Modeler",
  "Related Work 2.1 (Zhang fed his network into GeNIe).",
  "https://www.bayesfusion.com/genie/",
  "Product page for the GeNIe modeler. This is Zhang's ref [32]; cite "
  "as web reference with access date."),
 ("[30] Bureau of Transportation Statistics",
  "New methodology section, where departure counts feed the priors "
  "(Zhang's Eq. 6 uses BTS departures).",
  "https://www.transtats.bts.gov/",
  "Zhang's ref [45] is \"U.S. Air Carrier Aircraft Departures from 1975 "
  "to 2018\" from this site; his Table 6 reports those numbers. Cite "
  "the same source for your reproduction."),
 ("[31] Brier GW. Verification of forecasts expressed in terms of "
  "probability. Mon Weather Rev 1950;78(1):1-3.",
  "New held-out validation section, where the Brier score is first "
  "used.",
  "https://doi.org/10.1175/1520-0493(1950)078<0001:VOFEIT>2.0.CO;2",
  "Three-page paper defining the quadratic probability score (what we "
  "compute as \"Brier score\"). Free PDF via the AMS journal page."),
 ("[32] McNemar Q. Note on the sampling error of the difference between "
  "correlated proportions or percentages. Psychometrika "
  "1947;12(2):153-7.",
  "New held-out validation section, the paired significance tests; also "
  "old Section 9.4 already uses McNemar.",
  "https://doi.org/10.1007/BF02295996",
  "Defines the test on discordant pairs used for your paired "
  "predictor-vs-predictor comparisons."),
]

for title, where, link, says in REFS:
    p = doc.add_paragraph()
    p.add_run(title).bold = True
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.add_run("Goes where: ").bold = True
    p.add_run(where)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.add_run("Link: ").bold = True
    r = p.add_run(link)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x9D)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.add_run("What it says / how to verify: ").bold = True
    p.add_run(says)

doc.add_heading("Workflow reminder", level=1)
doc.add_paragraph(
    "Rewrite each item in your own words (never paste AI prose), send the "
    "finished section back for fact-check and approval, then move on. "
    "Abstract gets rewritten LAST. Target: sections 1-5 fixed by Wednesday, "
    "new core sections drafted Thursday, full assembly + references "
    "Thursday night, in-person walkthrough with Maha on Friday.")

doc.save(OUT)
print(f"wrote {OUT} with {n} items + {len(GLOBAL)} global fixes")
