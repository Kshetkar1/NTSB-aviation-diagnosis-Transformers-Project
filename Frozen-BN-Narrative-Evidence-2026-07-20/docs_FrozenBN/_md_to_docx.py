"""Convert paper_writing_guide.md to paper_writing_guide.docx.

Lightweight markdown -> docx converter tailored to this single file.
Handles: H1-H4, bullets, tables, **bold**, `code`, horizontal rules, blank lines.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

SRC = Path(__file__).parent / "paper_writing_guide.md"
DST = Path(__file__).parent / "paper_writing_guide.docx"

INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_CODE = re.compile(r"`([^`]+)`")
INLINE_ITALIC = re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)")


def add_runs(paragraph, text: str) -> None:
    """Render inline **bold**, *italic*, and `code` markers as runs."""

    pieces: list[tuple[str, set[str]]] = [(text, set())]

    def apply_pattern(pattern, marker: str) -> None:
        nonlocal pieces
        new_pieces: list[tuple[str, set[str]]] = []
        for chunk, marks in pieces:
            if marker in marks:
                new_pieces.append((chunk, marks))
                continue
            last_end = 0
            for m in pattern.finditer(chunk):
                if m.start() > last_end:
                    new_pieces.append((chunk[last_end : m.start()], set(marks)))
                new_pieces.append((m.group(1), marks | {marker}))
                last_end = m.end()
            if last_end < len(chunk):
                new_pieces.append((chunk[last_end:], set(marks)))
        pieces = new_pieces

    apply_pattern(INLINE_BOLD, "bold")
    apply_pattern(INLINE_CODE, "code")
    apply_pattern(INLINE_ITALIC, "italic")

    for chunk, marks in pieces:
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        if "bold" in marks:
            run.bold = True
        if "italic" in marks:
            run.italic = True
        if "code" in marks:
            run.font.name = "Consolas"
            run.font.size = Pt(10)


def setup(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def md_to_docx() -> None:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup(doc)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.rstrip()

        if not stripped.strip():
            i += 1
            continue

        if stripped.strip() == "---":
            doc.add_paragraph().add_run("_" * 60)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2), level=level)
            i += 1
            continue

        if stripped.lstrip().startswith("|") and i + 1 < len(lines) and re.match(
            r"^\s*\|[\s\-:|]+\|\s*$", lines[i + 1]
        ):
            header_cells = [c.strip() for c in stripped.strip().strip("|").split("|")]
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                row_cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(row_cells)
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header_cells))
            table.style = "Table Grid"
            for c_idx, txt in enumerate(header_cells):
                cell = table.rows[0].cells[c_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(txt)
                run.bold = True
            for r_idx, row in enumerate(rows, start=1):
                for c_idx in range(len(header_cells)):
                    txt = row[c_idx] if c_idx < len(row) else ""
                    cell = table.rows[r_idx].cells[c_idx]
                    cell.text = ""
                    add_runs(cell.paragraphs[0], txt)
            continue

        if stripped.lstrip().startswith("- "):
            content = stripped.lstrip()[2:]
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, content)
            i += 1
            continue

        num_match = re.match(r"^(\d+)\.\s+(.*)$", stripped.lstrip())
        if num_match:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, num_match.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(DST)
    print(f"Wrote {DST}")


if __name__ == "__main__":
    md_to_docx()
