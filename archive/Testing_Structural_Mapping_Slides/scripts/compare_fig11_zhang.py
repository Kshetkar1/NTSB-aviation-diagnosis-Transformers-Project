#!/usr/bin/env python3
"""
Compare our model against Zhang Fig 11: Main Gear Collapse Outcomes.

Zhang shows: given different gear failure evidence, what are the probabilities
of destroyed aircraft, minor damage, and minor injury?

All prognosis — uses embedding-based event matching.
Runs both A0 and A2. Outputs formatted .xlsx.

Usage:
    python compare_fig11_zhang.py
"""
from __future__ import annotations

import os
os.environ["NTSB_USE_TRAIN_INDEX"] = "1"

import json
import sys
from pathlib import Path

_SCRIPTS = Path(__file__).resolve().parent
PROJECT_ROOT = _SCRIPTS.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(_SCRIPTS))

import main_app  # noqa: E402
from config import LLM_MODEL, get_openai_api_key  # noqa: E402
from openai import OpenAI  # noqa: E402

from cause_code_mapper import match_events_to_targets  # noqa: E402
from paths import EVAL_OUTPUT_DIR  # noqa: E402

from extract_struct_v2 import extract_struct_v2  # noqa: E402
from io_cache import load_struct_jsonl  # noqa: E402
from paths import DEFAULT_STRUCT_CACHE_V2_PATH  # noqa: E402
from struct_hooks import make_score_adjust_fn  # noqa: E402
from struct_score_v2 import structural_similarity as structural_similarity_v2  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("WARNING: openpyxl not installed. Will output JSON only. pip install openpyxl")


EVIDENCE_LABELS = [
    "Landing main gear strut failure",
    "Landing gear emergency extension assembly failure",
    "Landing gear locking mechanism failure",
    "Landing main gear attachment failure",
]

EVIDENCE_QUERIES = [
    "The landing main gear strut failed during the landing phase, causing the main gear to collapse.",
    "The landing gear emergency extension assembly failed, preventing the gear from extending properly.",
    "The landing gear locking mechanism failed during landing, and the gear did not lock in position.",
    "The landing main gear attachment failed, causing the main gear to separate during landing.",
]

ZHANG_OUTCOMES = {
    "Destroyed aircraft damage": [0.045, 0.065, 0.055, 0.065],
    "Minor aircraft damage":     [0.12,  0.35,  0.30,  0.33],
    "Minor personnel injury":    [0.08,  0.25,  0.24,  0.32],
}


def run_prognosis(query, score_adjust_fn=None):
    r = main_app.predict_future_events(
        query, top_n_incidents=50, max_chain_steps=1, score_adjust_fn=score_adjust_fn,
    )
    return r.get("future_events", [])


def build_a2_adjust_fn(query, struct_by_eid):
    q_struct = extract_struct_v2(query, "zhang_fig11_query", model=LLM_MODEL)
    return make_score_adjust_fn(
        q_struct, struct_by_eid, alpha=0.5,
        similarity_fn=structural_similarity_v2,
    )


def write_xlsx(output_dir, all_results):
    if not HAS_OPENPYXL:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Fig 11 Comparison"

    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    zhang_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    a0_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    a2_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    bold_font = Font(bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    ws.cell(row=1, column=1, value="Evidence Condition").font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    col = 2
    for outcome in ZHANG_OUTCOMES:
        for model in ["Zhang BN", "A0", "A2"]:
            cell = ws.cell(row=1, column=col, value=f"{model} — {outcome}")
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin_border
            col += 1

    for row_idx, r in enumerate(all_results, start=2):
        ws.cell(row=row_idx, column=1, value=r["evidence"]).font = bold_font
        ws.cell(row=row_idx, column=1).border = thin_border
        col = 2
        for outcome in ZHANG_OUTCOMES:
            for val, fill in [
                (r[f"{outcome}_zhang"], zhang_fill),
                (r[f"{outcome}_a0"], a0_fill),
                (r[f"{outcome}_a2"], a2_fill),
            ]:
                cell = ws.cell(row=row_idx, column=col, value=round(val, 6))
                cell.number_format = "0.0000"
                cell.fill = fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")
                col += 1

    ws.column_dimensions["A"].width = 45
    xlsx_path = output_dir / "comparison_fig11.xlsx"
    wb.save(xlsx_path)
    return xlsx_path


def main():
    if not main_app.DATA_LOADED:
        raise RuntimeError("Train index not loaded.")

    output_dir = EVAL_OUTPUT_DIR / "zhang_comparison"
    output_dir.mkdir(parents=True, exist_ok=True)

    print("Loading structural cache for A2...", flush=True)
    struct_by_eid = load_struct_jsonl(DEFAULT_STRUCT_CACHE_V2_PATH)

    outcome_targets = list(ZHANG_OUTCOMES.keys())

    print("\n" + "=" * 70)
    print("FIG 11: Main Gear Collapse — Outcome Probabilities")
    print("  All prognosis — Zhang BN vs A0 vs A2 — direct target matching")
    print("=" * 70)

    all_results = []

    for i, (ev_label, query) in enumerate(zip(EVIDENCE_LABELS, EVIDENCE_QUERIES)):
        print(f"\n{'='*50}")
        print(f"Evidence {i+1}: {ev_label}")
        print(f"{'='*50}")

        print("  A0 prognosis...", flush=True)
        a0_prog = run_prognosis(query)
        a0_matched = match_events_to_targets(a0_prog, outcome_targets)

        print("  A2 prognosis...", flush=True)
        a2_adj = build_a2_adjust_fn(query, struct_by_eid)
        a2_prog = run_prognosis(query, score_adjust_fn=a2_adj)
        a2_matched = match_events_to_targets(a2_prog, outcome_targets)

        result = {"evidence": ev_label, "query": query}

        for outcome, zhang_vals in ZHANG_OUTCOMES.items():
            zhang_p = zhang_vals[i]
            a0_p = a0_matched.get(outcome, 0.0)
            a2_p = a2_matched.get(outcome, 0.0)

            result[f"{outcome}_zhang"] = zhang_p
            result[f"{outcome}_a0"] = a0_p
            result[f"{outcome}_a2"] = a2_p

            print(f"  {outcome:35s}  Zhang={zhang_p:.4f}  A0={a0_p:.4f}  A2={a2_p:.4f}")

        result["a0_prog_raw_top10"] = [{"event": e["event"], "probability": round(e["probability"], 6)} for e in a0_prog[:10]]
        result["a2_prog_raw_top10"] = [{"event": e["event"], "probability": round(e["probability"], 6)} for e in a2_prog[:10]]
        all_results.append(result)

    xlsx_path = write_xlsx(output_dir, all_results)

    json_path = output_dir / "comparison_fig11.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(all_results, f, indent=2, ensure_ascii=False, default=str)

    print(f"\n{'='*70}")
    print(f"Fig 11 comparison complete.")
    if xlsx_path:
        print(f"  XLSX: {xlsx_path}")
    print(f"  JSON: {json_path}")
    print(f"{'='*70}")


if __name__ == "__main__":
    main()
